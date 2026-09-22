"""Kunstwerk-inspectierapport in de huisstijl: PDF en Excel.

Wat hier vastligt:

1. **Excel is hetzelfde rapport.** Tabbladen Samenvatting, Elementen, Gebreken
   en Checklist, met het briefhoofd van de klant en de kopregel op rij 11.
2. **PDF en Excel noemen een kolom hetzelfde.** De kolommen van de PDF-tabellen
   komen, in dezelfde volgorde, uit hetzelfde blad als het Excel-tabblad.
3. **Excel rekent.** Conditie, ernst en score zijn getallen, geen tekst; een
   ingebedde foto (base64) komt niet in een cel terecht.
4. **De PDF overleeft rommel.** Logo, onbereikbare foto's en tekens buiten de
   PDF-letter (→, emoji) breken het rapport niet.
5. **Een andere organisatie ziet niets**, ook niet via de exports.
"""

import base64
import io
import re

import pytest
from openpyxl import load_workbook
from PIL import Image, ImageDraw

import export_huisstijl as h
import routers.kunstwerken_inspecties_router as kir
from database import SessionLocal
from models import AccountStatus, Asset, InspectionDefect, Organization, SubscriptionPlan
from tests.conftest import _make_user, auth
from tests.test_kunstwerken_inspecties import _make_asset, _new_inspection

KLANT = "Van Beek Infra B.V."
BASIS = "/api/kunstwerken-inspecties"


# ─────────────────────────────────────────────────────────────────────────────
# Helpers
# ─────────────────────────────────────────────────────────────────────────────

def _data_url(beeld: Image.Image) -> str:
    buf = io.BytesIO()
    beeld.save(buf, "PNG")
    return "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode()


def _foto(tint=(150, 150, 146), scheur=True) -> str:
    """Een 'foto' van beton met een scheur — genoeg om te zien hoe hij valt."""
    b, hh = 400, 300
    beeld = Image.new("RGB", (b, hh), tint)
    teken = ImageDraw.Draw(beeld)
    for i in range(0, hh, 6):
        grijs = tint[0] - 20 + (i * 7) % 40
        teken.line([(0, i), (b, i + 3)], fill=(grijs, grijs, grijs - 4), width=3)
    if scheur:
        teken.line([(40, 30), (140, 120), (180, 150), (260, 230), (360, 280)],
                   fill=(40, 38, 36), width=5)
    return _data_url(beeld)


def _logo() -> str:
    beeld = Image.new("RGBA", (600, 180), (0, 0, 0, 0))
    teken = ImageDraw.Draw(beeld)
    teken.rounded_rectangle([0, 10, 160, 170], radius=24, fill=(22, 101, 52, 255))
    teken.polygon([(40, 140), (80, 40), (120, 140)], fill=(255, 255, 255, 255))
    teken.rectangle([190, 55, 580, 85], fill=(22, 101, 52, 255))
    teken.rectangle([190, 105, 470, 125], fill=(100, 116, 139, 255))
    return _data_url(beeld)


def _handtekening() -> str:
    beeld = Image.new("RGBA", (420, 140), (0, 0, 0, 0))
    ImageDraw.Draw(beeld).line([(20, 100), (70, 30), (110, 110), (160, 40), (210, 95),
                                (260, 60), (330, 90), (400, 50)], fill=(15, 23, 42, 255), width=4)
    return _data_url(beeld)


def _klant_instellen(user, *, logo=True):
    db = SessionLocal()
    try:
        o = db.query(Organization).filter(Organization.id == user.organization_id).first()
        o.name = KLANT
        o.logo_data_url = _logo() if logo else None
        o.billing_address = "Industrieweg 12\n3044 AS Rotterdam"
        o.kvk_number = "24123456"
        o.btw_number = "NL812345678B01"
        o.contact_email = "inspecties@vanbeek.nl"
        o.contact_phone = "010 - 123 45 67"
        db.commit()
    finally:
        db.close()


def _defect(client, user, insp_id, el_id, **velden):
    r = client.post(f"{BASIS}/{insp_id}/elementen/{el_id}/defecten",
                    json=velden, headers=auth(user))
    assert r.status_code == 200, r.text
    return r.json()["id"]


def _antwoord(client, user, insp_id, el_id, code, **velden):
    r = client.patch(f"{BASIS}/{insp_id}/elementen/{el_id}/vragen/{code}",
                     json=velden, headers=auth(user))
    assert r.status_code == 200, r.text


def _realistische_inspectie(client, user, *, titel="Hoofdinspectie 2026",
                            omschrijving_extra="", https_foto=False, ondertekenen=True):
    """Een brug met meerdere bouwdelen, gebreken van licht tot kritiek, foto's,
    checklist-antwoorden en een handtekening."""
    db = SessionLocal()
    try:
        a = _make_asset(db, user=user, asset_type="brug", code="KW-0417")
        obj = db.query(Asset).filter(Asset.id == a.id).first()
        obj.name = "Brug over de Vaart"
        obj.properties_json = ('{"bouwjaar": "1978", "beheerder": "Gemeente Testdam", '
                               '"wegnummer": "N201"}')
        obj.lat, obj.lng = 52.160114, 4.497010
        obj.location_description = "Kruising Vaartweg / Dorpsstraat, noordzijde"
        db.commit()
        a_id = obj.id
    finally:
        db.close()

    insp = _new_inspection(
        client, user, a_id, title=titel, inspectie_type="Hoofdinspectie",
        datum_inspectie="2026-09-14T09:30:00", inspecteur_naam="J. de Vries",
        inspecteur_certificaat="CROW-KW 2024-117", weersomstandigheden="Droog, 14 °C, bewolkt",
        opdrachtgever_naam="Gemeente Testdam", opdrachtgever_email="beheer@testdam.nl")
    i_id = insp["id"]
    el = {e["element_code"]: e["id"] for e in insp["elementen"]}

    d1 = _defect(client, user, i_id, el["BRUG.ONDERBOUW"], gebrek_naam="Scheurvorming in beton",
                 gebrek_code="scheurvorming", ernst=2, intensiteit=2, omvang_klasse=3,
                 omvang_percentage=12.5, locatie_beschrijving="Landhoofd west, onderzijde",
                 omschrijving="Scheuren tot 0,4 mm breed over circa 2 m" + omschrijving_extra,
                 photo_url=_foto(), photo_url_2=_foto(tint=(170, 165, 150)))
    _defect(client, user, i_id, el["BRUG.ONDERBOUW"],
            gebrek_naam="Vochtdoorslag / kalkuitbloeiing", gebrek_code="vochtdoorslag",
            ernst=1, intensiteit=1, omvang_klasse=2, locatie_beschrijving="Landhoofd oost")
    _defect(client, user, i_id, el["BRUG.BOVENBOUW"], gebrek_naam="Corrosie staalconstructie",
            gebrek_code="corrosie_staal", ernst=3, intensiteit=3, omvang_klasse=4,
            locatie_beschrijving="Hoofdligger 2, veld 3",
            omschrijving="Doorroesting van de onderflens; plaatselijk materiaalverlies",
            gw_maatregel="Stralen, conserveren en plaatselijk versterken",
            photo_url=_foto(tint=(120, 90, 70), scheur=False))
    _defect(client, user, i_id, el["BRUG.BRUGDEK"], gebrek_naam="Scheurvorming in asfalt",
            gebrek_code="scheurvorming-langs", ernst=2, intensiteit=2, omvang_klasse=2,
            locatie_beschrijving="Rijstrook noord", crow_klasse="M2")
    d5 = _defect(client, user, i_id, el["BRUG.LEUNINGEN"], gebrek_naam="Aanrijdschade / deuk",
                 gebrek_code="deuk", ernst=2, intensiteit=3, omvang_klasse=1,
                 locatie_beschrijving="Leuning zuidzijde, paal 7-8")
    if https_foto:
        db = SessionLocal()
        try:
            dd = db.query(InspectionDefect).filter(InspectionDefect.id == d5).first()
            dd.photo_url = "https://nonexistent.invalid/defect.jpg"  # onbereikbaar
            db.commit()
        finally:
            db.close()
    assert d1

    _antwoord(client, user, i_id, el["BRUG.ONDERBOUW"], "GEN.STAAT", answer_score=4,
              toelichting="Scheurvorming aan de westzijde")
    _antwoord(client, user, i_id, el["BRUG.ONDERBOUW"], "GEN.VEILIG", answer_bool=False)
    _antwoord(client, user, i_id, el["BRUG.ONDERBOUW"], "GEN.INSPECTEERBAARHEID",
              answer_value_text="gedeeltelijk", toelichting="Waterzijde alleen vanaf de kant")
    _antwoord(client, user, i_id, el["BRUG.BOVENBOUW"], "GEN.STAAT", answer_score=5)
    _antwoord(client, user, i_id, el["BRUG.BOVENBOUW"], "GEN.VEILIG", answer_bool=True,
              toelichting="Draagvermogen onderflens twijfelachtig")

    for code, velden in (
        ("BRUG.ONDERBOUW", {"bevindingen": "Beton oogt verder gaaf; scheuren niet actief.",
                            "aanbevolen_actie": "Scheuren injecteren binnen 2 jaar"}),
        ("BRUG.BOVENBOUW", {"bevindingen": "Corrosie concentreert zich bij de lekkende voeg.",
                            "aanbevolen_actie": "Constructieve beoordeling laten uitvoeren"}),
        ("BRUG.PIJLERS", {"beoordeeld": True}),
        ("BRUG.OPLEGGINGEN", {"niet_inspecteerbaar_reden": "Niet bereikbaar zonder hoogwerker"}),
        ("BRUG.AFWATERING", {"beoordeeld": True}),
    ):
        r = client.patch(f"{BASIS}/{i_id}/elementen/{el[code]}", json=velden, headers=auth(user))
        assert r.status_code == 200, r.text

    r = client.patch(f"{BASIS}/{i_id}", json={
        "samenvatting": "De brug is functioneel, maar de hoofdligger vraagt op korte termijn aandacht",
        "aanbevolen_acties": "Hoofdligger 2 conserveren; voegovergang vervangen",
        "bijzonderheden": "Tijdens de inspectie was rijstrook noord afgezet (verkeersmaatregel).",
    }, headers=auth(user))
    assert r.status_code == 200, r.text
    if ondertekenen:
        r = client.post(f"{BASIS}/{i_id}/complete", headers=auth(user))
        assert r.status_code == 200, r.text
        r = client.post(f"{BASIS}/{i_id}/sign", json={"signature_data_url": _handtekening()},
                        headers=auth(user))
        assert r.status_code == 200, r.text
    return i_id


def _werkboek(r):
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith(h.XLSX_MIME)
    return load_workbook(io.BytesIO(r.content))


def _kopregel(ws):
    return [c.value for c in ws[h.KOP_RIJ_EXCEL] if c.value is not None]


def _is_deelrij(deel, geheel) -> bool:
    """Staan alle namen van `deel` in `geheel`, in dezelfde volgorde?"""
    it = iter(geheel)
    return all(naam in it for naam in deel)


class _OnverpaktePDF(h.HuisstijlPDF):
    """Zonder compressie, zodat de tekst in de PDF te doorzoeken is."""

    def __init__(self, *a, **kw):
        super().__init__(*a, **kw)
        self.compress = False


# ─────────────────────────────────────────────────────────────────────────────
# Excel
# ─────────────────────────────────────────────────────────────────────────────

def test_xlsx_rapport_in_huisstijl(client, admin_user):
    _klant_instellen(admin_user)
    i_id = _realistische_inspectie(client, admin_user)
    r = client.get(f"{BASIS}/{i_id}/export.xlsx", headers=auth(admin_user))
    wb = _werkboek(r)
    assert re.search(r'filename="inspectierapport-KW-0417-\d{4}-\d{2}-\d{2}\.xlsx"',
                     r.headers["content-disposition"])
    assert wb.sheetnames == ["Samenvatting", "Elementen", "Gebreken", "Checklist"]

    for ws in wb.worksheets:
        briefhoofd = [c.value for rij in ws.iter_rows(min_row=1, max_row=4) for c in rij]
        assert KLANT in briefhoofd, ws.title
        assert "KvK 24123456 · BTW NL812345678B01" in briefhoofd, ws.title
        assert ws._images, f"logo ontbreekt op {ws.title}"

    # Kopregel op rij 11; de PDF-tabellen tonen dezelfde kolommen in dezelfde volgorde
    assert _kopregel(wb["Samenvatting"]) == ["Onderdeel", "Waarde"]
    assert _kopregel(wb["Elementen"])[:len(kir._PDF_ELEMENTEN)] == list(kir._PDF_ELEMENTEN)
    assert _kopregel(wb["Gebreken"])[:len(kir._PDF_GEBREKEN)] == list(kir._PDF_GEBREKEN)
    assert _is_deelrij(kir._PDF_TOP_GEBREKEN, _kopregel(wb["Gebreken"]))
    assert _is_deelrij(kir._PDF_CHECKLIST, _kopregel(wb["Checklist"]))


def test_xlsx_bevat_dezelfde_gegevens_als_de_csv_met_echte_getallen(client, admin_user):
    i_id = _realistische_inspectie(client, admin_user)
    wb = _werkboek(client.get(f"{BASIS}/{i_id}/export.xlsx", headers=auth(admin_user)))
    csv_tekst = client.get(f"{BASIS}/{i_id}/export.csv", headers=auth(admin_user)).text
    csv_regels = csv_tekst.splitlines()
    n_elementen = csv_regels.index("=== DEFECTEN ===") - csv_regels.index("=== ELEMENTEN ===") - 3
    detail = client.get(f"{BASIS}/{i_id}", headers=auth(admin_user)).json()

    el = wb["Elementen"]
    kop = _kopregel(el)
    rijen = [dict(zip(kop, r)) for r in el.iter_rows(min_row=h.KOP_RIJ_EXCEL + 1,
                                                       max_col=len(kop), values_only=True)]
    rijen = [r for r in rijen if r["Bouwdeelcode"]]      # zonder de toelichting eronder
    assert len(rijen) == n_elementen == len(detail["elementen"])
    per_code = {r["Bouwdeelcode"]: r for r in rijen}
    onderbouw = per_code["BRUG.ONDERBOUW"]
    assert onderbouw["Bouwdeel"] == "Onderbouw / landhoofden"
    assert isinstance(onderbouw["Conditie"], int) and onderbouw["Gebreken"] == 2
    assert onderbouw["Aanbevolen actie"] == "Scheuren injecteren binnen 2 jaar"

    gb = wb["Gebreken"]
    kop = _kopregel(gb)
    gebreken = [dict(zip(kop, r)) for r in gb.iter_rows(min_row=h.KOP_RIJ_EXCEL + 1,
                                                          max_col=len(kop), values_only=True)
                if r[1]]
    assert len(gebreken) == sum(len(e["defecten"]) for e in detail["elementen"]) == 5
    corrosie = next(g for g in gebreken if g["Gebrek"] == "Corrosie staalconstructie")
    assert isinstance(corrosie["Score"], int) and corrosie["Score"] >= 5
    assert corrosie["Ernst"] == 3 and corrosie["Bouwdeelcode"] == "BRUG.BOVENBOUW"
    scheur = next(g for g in gebreken if g["Gebrek"] == "Scheurvorming in beton")
    assert scheur["Omvang (%)"] == 12.5
    assert scheur["Foto"] == "in PDF-rapport"      # twee ingebedde foto's, één verwijzing
    asfalt = next(g for g in gebreken if g["Gebrek"] == "Scheurvorming in asfalt")
    assert asfalt["CROW-klasse"] == "M2" and "€" in asfalt["Maatregel"]

    # Geen base64-foto of handtekening in een cel (Excel kapt af bij 32.767 tekens)
    for ws in wb.worksheets:
        for rij in ws.iter_rows(values_only=True):
            for v in rij:
                if isinstance(v, str):
                    assert "data:image" not in v and len(v) < 32767

    samen = {r[0]: r[1] for r in wb["Samenvatting"].iter_rows(
        min_row=h.KOP_RIJ_EXCEL + 1, max_col=2, values_only=True) if r[0]}
    assert samen["Object"] == "KW-0417 – Brug over de Vaart"
    assert samen["Inspecteur"] == "J. de Vries (CROW-KW 2024-117)"
    assert samen["Status"] == "Ondertekend"
    assert samen["Eindconditie NEN 2767-2"].startswith(str(detail["conditiescore_overall"]))
    assert samen["Gebreken"].startswith("5, waarvan")

    checklist = wb["Checklist"]
    kop = _kopregel(checklist)
    antwoorden = [dict(zip(kop, r)) for r in checklist.iter_rows(
        min_row=h.KOP_RIJ_EXCEL + 1, max_col=len(kop), values_only=True) if r[1]]
    veilig = next(a for a in antwoorden if a["Bouwdeelcode"] == "BRUG.BOVENBOUW"
                  and a["Vraag"] == "Veiligheidsrisico geconstateerd?")
    assert veilig["Antwoord"] == "Ja" and veilig["Aandacht"] == "ja"


def test_xlsx_van_lege_draft(client, admin_user):
    db = SessionLocal()
    try:
        a_id = _make_asset(db, user=admin_user, asset_type="viaduct", code="KW-LEEG").id
    finally:
        db.close()
    insp = _new_inspection(client, admin_user, a_id, kunstwerk_type="viaduct", auto_elements=False)
    wb = _werkboek(client.get(f"{BASIS}/{insp['id']}/export.xlsx", headers=auth(admin_user)))
    assert wb.sheetnames == ["Samenvatting", "Elementen", "Gebreken", "Checklist"]
    assert wb["Gebreken"].cell(row=h.KOP_RIJ_EXCEL + 1, column=1).value == \
        "Geen gegevens voor deze selectie."
    # ook de PDF van een inspectie zonder bouwdelen
    r = client.get(f"{BASIS}/{insp['id']}/export.pdf", headers=auth(admin_user))
    assert r.status_code == 200, r.text
    assert r.content[:4] == b"%PDF"


# ─────────────────────────────────────────────────────────────────────────────
# PDF
# ─────────────────────────────────────────────────────────────────────────────

def test_pdf_rapport_in_huisstijl_met_logo_fotos_en_handtekening(client, admin_user, monkeypatch):
    monkeypatch.setattr(kir, "HuisstijlPDF", _OnverpaktePDF)
    _klant_instellen(admin_user)
    i_id = _realistische_inspectie(client, admin_user)
    r = client.get(f"{BASIS}/{i_id}/export.pdf", headers=auth(admin_user))
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/pdf")
    assert re.search(r'filename="inspectierapport-KW-0417-\d{4}-\d{2}-\d{2}\.pdf"',
                     r.headers["content-disposition"])
    ruw = r.content.decode("latin-1")
    assert ruw.startswith("%PDF")
    for verwacht in (KLANT, "KvK 24123456", "Inspectierapport", "Inhoudsopgave",
                     "4. Bevindingen per bouwdeel", "7. Verantwoording en ondertekening",
                     "Belangrijkste gebreken", "CHECKLIST NEN/CROW", "Aandachtspunten",
                     "Intensiteit", "Normreferentie", "Gemeente Testdam", "Pagina 1 van",
                     "Opgesteld met FieldOps"):
        assert verwacht in ruw, verwacht
    # logo + 3 defectfoto's + handtekening
    assert len(re.findall(r"/Subtype /Image", ruw)) >= 5
    # Inhoudsopgave met paginanummers en bladwijzers per hoofdstuk
    assert ruw.count("/Outlines") >= 1
    detail = client.get(f"{BASIS}/{i_id}", headers=auth(admin_user)).json()
    assert detail["pdf_generated_at"] is not None


def test_pdf_met_onbereikbare_foto_en_rare_tekens(client, admin_user):
    _klant_instellen(admin_user, logo=False)
    titel = "Brug “De Hoop” → hoofdinspectie 🚧 € 12.500 – ≥ 3 ✓"
    i_id = _realistische_inspectie(client, admin_user, titel=titel, https_foto=True,
                                   omschrijving_extra=" → verder 🙂 ≈ 2 m² ≠ oké",
                                   ondertekenen=False)
    r = client.get(f"{BASIS}/{i_id}/export.pdf", headers=auth(admin_user))
    assert r.status_code == 200, r.text
    assert r.content[:4] == b"%PDF"
    # Excel is Unicode: daar blijft de titel precies zoals ingevoerd
    wb = _werkboek(client.get(f"{BASIS}/{i_id}/export.xlsx", headers=auth(admin_user)))
    samen = {r[0]: r[1] for r in wb["Samenvatting"].iter_rows(
        min_row=h.KOP_RIJ_EXCEL + 1, max_col=2, values_only=True) if r[0]}
    assert samen["Titel inspectie"] == titel
    assert samen["Status"] != "Ondertekend"


# ─────────────────────────────────────────────────────────────────────────────
# Toegang
# ─────────────────────────────────────────────────────────────────────────────

@pytest.mark.parametrize("formaat", ["pdf", "xlsx"])
def test_andere_organisatie_kan_niet_exporteren(client, admin_user, formaat):
    db = SessionLocal()
    try:
        a_id = _make_asset(db, user=admin_user, asset_type="brug", code="KW-GEHEIM").id
        other_org = Organization(name="OtherOrg", plan=SubscriptionPlan.PROFESSIONAL,
                                 status=AccountStatus.ACTIVE, max_users=10)
        db.add(other_org)
        db.commit()
        db.refresh(other_org)
        other_user = _make_user(db, "other-export@test.nl", org=other_org)
    finally:
        db.close()
    insp = _new_inspection(client, admin_user, a_id)
    r = client.get(f"{BASIS}/{insp['id']}/export.{formaat}", headers=auth(other_user))
    assert r.status_code == 404
    r = client.get(f"{BASIS}/does-not-exist/export.{formaat}", headers=auth(admin_user))
    assert r.status_code == 404
