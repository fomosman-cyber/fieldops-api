"""Gebouwinspectie, toolbox en werkplekinspectie in de huisstijl: PDF en Excel.

Wat hier vastligt:

1. **Het briefhoofd is van de klant.** Logo, naam en gegevens komen uit de
   organisatie die exporteert, in PDF én Excel.
2. **Excel leest als de PDF.** Elk tabblad heeft zijn kopregel op rij 11, met
   precies de kolomnamen van de tabel in de PDF.
3. **Rare tekens breken niets.** Pijlen, emoji en het euroteken in wat een
   gebruiker typt leveren nog steeds een PDF op.
4. **Niemand exporteert andermans inspectie.** Een andere organisatie krijgt
   een 404, net als bij het openen van de inspectie zelf.
"""
import base64
import io
import re
import zlib

import pytest
from openpyxl import load_workbook
from PIL import Image, ImageDraw

import bouw_boei as bb
import wpi_checklist as wc
from database import SessionLocal
from export_huisstijl import KOP_RIJ_EXCEL, XLSX_MIME
from models import AccountStatus, Organization, Project, SubscriptionPlan
from routers import bouw_router, toolbox_router, wpi_router

from .conftest import _make_user, auth

KLANT = "Van Beek Infra B.V."
RAAR = "Gat → verzakking 🚧 € 1.250 – “spoed” ✓"


# ---------------------------------------------------------------------------
# Hulpjes
# ---------------------------------------------------------------------------

def _png(b=300, h=100, kleur=(22, 101, 52, 255)) -> str:
    buf = io.BytesIO()
    Image.new("RGBA", (b, h), kleur).save(buf, "PNG")
    return "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode()


def _handtekening() -> str:
    beeld = Image.new("RGBA", (600, 200), (0, 0, 0, 0))
    ImageDraw.Draw(beeld).line([(20, 150), (120, 40), (200, 160), (320, 50), (560, 120)],
                               fill=(20, 30, 90, 255), width=6)
    buf = io.BytesIO()
    beeld.save(buf, "PNG")
    return "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode()


def _foto() -> str:
    buf = io.BytesIO()
    Image.new("RGB", (800, 600), (120, 110, 100)).save(buf, "JPEG")
    return "data:image/jpeg;base64," + base64.b64encode(buf.getvalue()).decode()


@pytest.fixture
def klant(org):
    """De organisatie van de test krijgt een logo en volledige gegevens."""
    db = SessionLocal()
    try:
        o = db.query(Organization).filter(Organization.id == org.id).first()
        o.name = KLANT
        o.logo_data_url = _png()
        o.billing_address = "Industrieweg 12\n3044 AS Rotterdam"
        o.kvk_number = "24123456"
        o.btw_number = "NL812345678B01"
        o.contact_email = "planning@vanbeek.nl"
        o.contact_phone = "010 123 45 67"
        o.brand_color = "#ff00ff"       # hoort in de export niet terug te komen
        db.commit()
        return o.id
    finally:
        db.close()


@pytest.fixture
def andere_admin():
    db = SessionLocal()
    try:
        andere = Organization(name="Concurrent BV", plan=SubscriptionPlan.PROFESSIONAL,
                              status=AccountStatus.ACTIVE, max_users=10)
        db.add(andere)
        db.commit()
        db.refresh(andere)
        return _make_user(db, "concurrent@test.nl", org=andere)
    finally:
        db.close()


def _project(user, naam="N207 Alphen – fase 2") -> str:
    db = SessionLocal()
    try:
        p = Project(name=naam, organization_id=user.organization_id, status="active",
                    created_by=user.id)
        db.add(p)
        db.commit()
        return p.id
    finally:
        db.close()


def _pdf_tekst(inhoud: bytes) -> str:
    """Alle tekst uit de (gecomprimeerde) PDF-streams, om op te zoeken."""
    delen = []
    for stroom in re.findall(rb"stream\r?\n(.*?)\r?\nendstream", inhoud, re.S):
        try:
            delen.append(zlib.decompress(stroom).decode("latin-1"))
        except zlib.error:
            continue
    return "\n".join(delen)


def _aantal_afbeeldingen(inhoud: bytes) -> int:
    return len(re.findall(rb"/Subtype\s*/Image", inhoud))


def _werkboek(r):
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == XLSX_MIME
    assert r.headers["content-disposition"].endswith('.xlsx"')
    return load_workbook(io.BytesIO(r.content))


def _briefhoofd(ws) -> list:
    return [c.value for rij in ws.iter_rows(min_row=1, max_row=4) for c in rij if c.value]


def _kopregel(ws) -> list:
    return [c.value for c in ws[KOP_RIJ_EXCEL] if c.value is not None]


def _namen(kolommen) -> list:
    return [k.naam for k in kolommen]


# ---------------------------------------------------------------------------
# Gebouwinspectie (BOEI)
# ---------------------------------------------------------------------------

def _gebouw(client, user, naam="Gemeentehuis") -> str:
    r = client.post("/api/bouw/", headers=auth(user), json={
        "gebouw_naam": naam, "straatnaam": "Coolsingel", "huisnummer": "40",
        "postcode": "3011AD", "plaats": "Rotterdam", "gebouw_type": "kantoor",
        "pijlers": ["B", "I"]})
    assert r.status_code == 200, r.text
    bouw_id = r.json()["id"]
    antwoorden = r.json()["antwoorden"]
    for a in antwoorden:
        client.patch(f"/api/bouw/{bouw_id}/antwoorden/{a['id']}", headers=auth(user),
                     json={"antwoord": "ja"})
    client.patch(f"/api/bouw/{bouw_id}/antwoorden/{antwoorden[0]['id']}", headers=auth(user),
                 json={"antwoord": "nee", "toelichting": RAAR, "actie": "Vluchtroute vrijmaken"})
    # Een vraag met een bewijsstuk dat er niet is: komt bij de ontbrekende documenten.
    met_bewijs = {s["code"] for s in bb.bewijsstukken()}
    stuk = next(a for a in antwoorden if a["question_code"] in met_bewijs)
    client.patch(f"/api/bouw/{bouw_id}/antwoorden/{stuk['id']}", headers=auth(user),
                 json={"bewijs_aanwezig": False})
    client.post(f"/api/bouw/{bouw_id}/condities", headers=auth(user),
                json={"element_code": "DAK.01", "gebrek": "blaasvorming → lekkage 🌧",
                      "ernst": 3, "intensiteit": 2, "omvang_klasse": 4})
    return bouw_id


def test_bouw_excel_met_briefhoofd_en_de_kolommen_van_de_pdf(client, admin_user, klant):
    bouw_id = _gebouw(client, admin_user)
    wb = _werkboek(client.get(f"/api/bouw/{bouw_id}/export.xlsx", headers=auth(admin_user)))
    assert "Gemeentehuis" in client.get(f"/api/bouw/{bouw_id}/export.xlsx",
                                        headers=auth(admin_user)).headers["content-disposition"]

    verwacht = {
        "Samenvatting": bouw_router.KOLOMMEN_GEGEVENS,
        "Per pijler": bouw_router.KOLOMMEN_PIJLERS,
        "Conditie NEN 2767": bouw_router.KOLOMMEN_CONDITIE,
        "Aandachtspunten": bouw_router.KOLOMMEN_PUNTEN,
        "Ontbrekende documenten": bouw_router.KOLOMMEN_DOCUMENTEN,
    }
    assert wb.sheetnames == list(verwacht)
    for naam, kolommen in verwacht.items():
        ws = wb[naam]
        assert KLANT in _briefhoofd(ws), naam
        assert "KvK 24123456 · BTW NL812345678B01" in _briefhoofd(ws), naam
        assert _kopregel(ws) == _namen(kolommen), naam
    assert len(wb["Samenvatting"]._images) == 1          # het logo van de klant

    # De PDF heeft precies dezelfde kolomkoppen.
    pdf = client.get(f"/api/bouw/{bouw_id}/export.pdf", headers=auth(admin_user))
    tekst = _pdf_tekst(pdf.content)
    for kolommen in list(verwacht.values())[1:]:
        for k in kolommen:
            assert f"({k.naam})" in tekst, k.naam

    conditie = [c.value for c in wb["Conditie NEN 2767"][KOP_RIJ_EXCEL + 1]]
    assert conditie[0] == "DAK.01" and isinstance(conditie[6], int)
    punt = [c.value for c in wb["Aandachtspunten"][KOP_RIJ_EXCEL + 1]]
    assert punt[-1] == "open" and punt[4] == "Vluchtroute vrijmaken"
    assert wb["Ontbrekende documenten"].cell(KOP_RIJ_EXCEL + 1, 1).value
    samenvatting = {r[0].value: r[1].value for r in wb["Samenvatting"].iter_rows(min_row=12)
                    if r[0].value}
    assert samenvatting["Gebouw"] == "Gemeentehuis"


def test_bouw_pdf_met_logo_en_rare_tekens(client, admin_user, klant):
    bouw_id = _gebouw(client, admin_user, naam="Stadhuis 🏛 → west")
    r = client.get(f"/api/bouw/{bouw_id}/export.pdf", headers=auth(admin_user))
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == "application/pdf"
    assert r.content.startswith(b"%PDF")
    assert _aantal_afbeeldingen(r.content) >= 1            # het logo
    tekst = _pdf_tekst(r.content)
    assert KLANT in tekst and "KvK 24123456" in tekst
    assert "Opgesteld met FieldOps" in tekst
    # Header blijft ASCII, anders loopt de download stuk.
    r.headers["content-disposition"].encode("ascii")


def test_lange_toelichting_blijft_heel_en_leesbaar_in_excel(client, admin_user, klant):
    """De PDF breekt 'ontruimingsalarminstallatie' af met zachte streepjes;
    Excel krijgt de tekst zoals hij is, en een rij met een lange toelichting
    wordt hoog genoeg om hem te lezen."""
    bouw_id = _gebouw(client, admin_user)
    d = client.get(f"/api/bouw/{bouw_id}", headers=auth(admin_user)).json()
    lang = "Onderhoudscertificaat ontruimingsalarminstallatie verlopen, opnieuw laten keuren. " * 4
    client.patch(f"/api/bouw/{bouw_id}/antwoorden/{d['antwoorden'][0]['id']}",
                 headers=auth(admin_user), json={"antwoord": "nee", "toelichting": lang})
    wb = _werkboek(client.get(f"/api/bouw/{bouw_id}/export.xlsx", headers=auth(admin_user)))
    for ws in wb.worksheets:
        for rij in ws.iter_rows(values_only=True):
            assert not any(isinstance(v, str) and "­" in v for v in rij), ws.title
    ws = wb["Aandachtspunten"]
    assert ws.cell(KOP_RIJ_EXCEL + 1, 4).value == lang
    assert ws.row_dimensions[KOP_RIJ_EXCEL + 1].height > 30


def test_bouw_export_van_andere_organisatie_geeft_404(client, admin_user, andere_admin):
    bouw_id = _gebouw(client, admin_user)
    for ext in ("pdf", "xlsx"):
        assert client.get(f"/api/bouw/{bouw_id}/export.{ext}",
                          headers=auth(andere_admin)).status_code == 404


# ---------------------------------------------------------------------------
# Toolbox
# ---------------------------------------------------------------------------

def _toolbox(client, user) -> str:
    r = client.post("/api/toolbox/", headers=auth(user), json={
        "project_id": _project(user),
        "onderwerp": "Werken langs de rijbaan → " + RAAR,
        "inleiding": "Vandaag werken we langs de rijbaan 🚧 met € 500 boete-risico.",
        "risicos": ["Passerend verkeer – zeer dichtbij"],
        "maatregelen": ["Afzetting volgens CROW 96b ✓"],
        "bespreekpunten": ["Wie bewaakt de afzetting?"],
        "afspraken": "Om 07:00 verzamelen bij de keet.",
    })
    assert r.status_code == 200, r.text
    tid = r.json()["id"]
    for naam, bedrijf, teken in (("Jan de Vries", "De Vries Infra BV", True),
                                 ("Piet Jansen", None, False)):
        d = client.post(f"/api/toolbox/{tid}/deelnemers", headers=auth(user),
                        json={"naam": naam, "bedrijf": bedrijf}).json()
        if teken:
            client.post(f"/api/toolbox/{tid}/deelnemers/{d['id']}/sign", headers=auth(user),
                        json={"signature_data_url": _handtekening()})
    return tid


def test_toolbox_excel_met_briefhoofd_en_presentielijst(client, admin_user, klant):
    tid = _toolbox(client, admin_user)
    wb = _werkboek(client.get(f"/api/toolbox/{tid}/export.xlsx", headers=auth(admin_user)))
    assert wb.sheetnames == ["Samenvatting", "Aanwezigen"]
    for ws in wb.worksheets:
        assert KLANT in _briefhoofd(ws)
    assert _kopregel(wb["Samenvatting"]) == _namen(toolbox_router.KOLOMMEN_GEGEVENS)
    assert _kopregel(wb["Aanwezigen"]) == _namen(toolbox_router.KOLOMMEN_AANWEZIGEN)

    jan = [c.value for c in wb["Aanwezigen"][KOP_RIJ_EXCEL + 1]]
    assert jan[:3] == ["Jan de Vries", "De Vries Infra BV", "ja"]
    assert jan[3].startswith("getekend op ")
    piet = [c.value for c in wb["Aanwezigen"][KOP_RIJ_EXCEL + 2]]
    assert piet[0] == "Piet Jansen" and piet[3] is None

    onderdelen = [r[0].value for r in wb["Samenvatting"].iter_rows(min_row=KOP_RIJ_EXCEL + 1)]
    assert {"Project", "Risico 1", "Maatregel 1", "Bespreekpunt 1", "Afspraken"} <= set(onderdelen)

    pdf = client.get(f"/api/toolbox/{tid}/export.pdf", headers=auth(admin_user))
    tekst = _pdf_tekst(pdf.content)
    for k in toolbox_router.KOLOMMEN_AANWEZIGEN:
        assert f"({k.naam})" in tekst, k.naam


def test_toolbox_pdf_met_logo_handtekening_en_rare_tekens(client, admin_user, klant):
    tid = _toolbox(client, admin_user)
    r = client.get(f"/api/toolbox/{tid}/export.pdf", headers=auth(admin_user))
    assert r.status_code == 200, r.text
    assert r.content[:4] == b"%PDF"
    assert _aantal_afbeeldingen(r.content) >= 2            # logo + handtekening
    assert KLANT in _pdf_tekst(r.content)
    r.headers["content-disposition"].encode("ascii")


def test_toolbox_onleesbare_handtekening_breekt_de_pdf_niet(client, admin_user, klant):
    tid = _toolbox(client, admin_user)
    d = client.post(f"/api/toolbox/{tid}/deelnemers", headers=auth(admin_user),
                    json={"naam": "Kees"}).json()
    client.post(f"/api/toolbox/{tid}/deelnemers/{d['id']}/sign", headers=auth(admin_user),
                json={"signature_data_url": "data:image/png;base64,bm9vaXQ="})
    r = client.get(f"/api/toolbox/{tid}/export.pdf", headers=auth(admin_user))
    assert r.status_code == 200, r.text
    assert "(getekend)" in _pdf_tekst(r.content)


def test_toolbox_export_van_andere_organisatie_geeft_404(client, admin_user, andere_admin):
    tid = _toolbox(client, admin_user)
    for ext in ("pdf", "xlsx"):
        assert client.get(f"/api/toolbox/{tid}/export.{ext}",
                          headers=auth(andere_admin)).status_code == 404


# ---------------------------------------------------------------------------
# Werkplekinspectie
# ---------------------------------------------------------------------------

def _wpi(client, user) -> str:
    r = client.post("/api/wpi/", headers=auth(user),
                    json={"project_id": _project(user), "locatie": "Bij de inrit → oost"})
    assert r.status_code == 200, r.text
    w = r.json()
    for i, a in enumerate(w["antwoorden"]):
        body = {"antwoord": "nvt" if i % 5 == 4 else "ja"}
        if i == 0:
            body = {"antwoord": "nee", "toelichting": RAAR, "actie": "Plan ophangen in de keet",
                    "actiehouder_id": user.id, "photo_url": _foto()}
        client.patch(f"/api/wpi/{w['id']}/antwoorden/{a['id']}", headers=auth(user), json=body)
    client.patch(f"/api/wpi/{w['id']}", headers=auth(user),
                 json={"algemene_indruk": "Nette bouwplaats 👍, wel € 0 budget voor hekwerk."})
    return w["id"]


def test_wpi_excel_met_briefhoofd_acties_en_checklist(client, admin_user, klant):
    wid = _wpi(client, admin_user)
    wb = _werkboek(client.get(f"/api/wpi/{wid}/export.xlsx", headers=auth(admin_user)))
    assert wb.sheetnames == ["Samenvatting", "Acties", "Checklist"]
    for ws in wb.worksheets:
        assert KLANT in _briefhoofd(ws)
    assert _kopregel(wb["Samenvatting"]) == _namen(wpi_router.KOLOMMEN_GEGEVENS)
    assert _kopregel(wb["Acties"]) == _namen(wpi_router.KOLOMMEN_ACTIES)
    assert _kopregel(wb["Checklist"]) == _namen(wpi_router.KOLOMMEN_CHECKLIST)

    actie = [c.value for c in wb["Acties"][KOP_RIJ_EXCEL + 1]]
    assert actie[3] == "Plan ophangen in de keet" and actie[4] and actie[5] == "open"
    checklist = list(wb["Checklist"].iter_rows(min_row=KOP_RIJ_EXCEL + 1, values_only=True))
    assert len(checklist) == len(wc.VRAGEN)
    assert all(r[0] for r in checklist), "in Excel staat de categorie op elke regel"
    assert checklist[0][2] == "NIET IN ORDE"

    pdf = client.get(f"/api/wpi/{wid}/export.pdf", headers=auth(admin_user))
    tekst = _pdf_tekst(pdf.content)
    for k in (*wpi_router.KOLOMMEN_ACTIES, *wpi_router.KOLOMMEN_CHECKLIST):
        assert f"({k.naam})" in tekst, k.naam


def test_wpi_pdf_met_logo_foto_en_rare_tekens(client, admin_user, klant):
    wid = _wpi(client, admin_user)
    r = client.get(f"/api/wpi/{wid}/export.pdf", headers=auth(admin_user))
    assert r.status_code == 200, r.text
    assert r.content[:4] == b"%PDF"
    assert _aantal_afbeeldingen(r.content) >= 2            # logo + foto
    tekst = _pdf_tekst(r.content)
    assert KLANT in tekst and "Werkplekinspectie" in tekst


def test_wpi_export_van_andere_organisatie_geeft_404(client, admin_user, andere_admin):
    wid = _wpi(client, admin_user)
    for ext in ("pdf", "xlsx"):
        assert client.get(f"/api/wpi/{wid}/export.{ext}",
                          headers=auth(andere_admin)).status_code == 404
