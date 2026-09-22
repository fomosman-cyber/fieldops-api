"""MJOP-export als PDF én Excel, in één huisstijl.

Wat hier vastligt:

1. **Excel is de PDF, als werkboek.** Zelfde briefhoofd (naam en gegevens
   van de klant), en per tabel uit het rapport een tabblad met precies de
   kolommen die de PDF-tabel heeft. De kopregel staat altijd op rij 11.
2. **Excel rekent.** Bedragen zijn getallen met een euro-opmaak, jaartallen
   getallen zonder duizendtalpunt -- geen tekst.
3. **Geen verzonnen getallen.** Zonder indexpercentage geen kolommen met
   geïndexeerde bedragen; met een percentage wel.
4. **Klantdata breekt niets.** Een logo, pijlen, emoji en eurotekens in naam
   of project leveren gewoon een PDF en een werkboek op.
5. **Wat van een ander is, blijft van een ander.** Een gebruiker van een
   andere organisatie krijgt een export van zijn eigen (lege) MJOP, ook met
   het project-id van iemand anders in de url.
"""

import base64
import io
import json
import re
import zlib
from datetime import datetime, timedelta, timezone

import pytest
from openpyxl import load_workbook
from PIL import Image

import mjop_kosten as mjop
from database import SessionLocal
from models import (AccountStatus, Asset, Melding, Organization, Project,
                    SubscriptionPlan, UserRole)
from routers.mjop_router import _kolommen
from tests.conftest import _make_user, auth

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
KOP_RIJ = 11
BLADEN = ["Kosten per jaar", "MJOP-tabel", "Assets per type", "Conditie-verdeling",
          "Meldingen per categorie", "Assets met meldingen"]


# ─────────────────────────────────────────────────────────────────────────────
# Helpers
# ─────────────────────────────────────────────────────────────────────────────

def _logo_data_url():
    buf = io.BytesIO()
    Image.new("RGB", (300, 100), (22, 101, 52)).save(buf, "PNG")
    return "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode()


def _org_gegevens(org_id, **velden):
    db = SessionLocal()
    try:
        o = db.query(Organization).filter(Organization.id == org_id).first()
        for k, v in velden.items():
            setattr(o, k, v)
        db.commit()
    finally:
        db.close()


def _project(org_id, user_id, naam="Kunstwerken Noord"):
    db = SessionLocal()
    try:
        p = Project(name=naam, organization_id=org_id, status="active", created_by=user_id)
        db.add(p)
        db.commit()
        db.refresh(p)
        return p.id
    finally:
        db.close()


def _asset(org_id, user_id, *, code, asset_type="brug", score=4, jaren_vooruit=2,
           project_id=None, naam=None, lengte=None):
    db = SessionLocal()
    try:
        a = Asset(organization_id=org_id, created_by=user_id, code=code,
                  name=naam or f"Asset {code}", asset_type=asset_type,
                  condition_score=score, project_id=project_id, length_m=lengte,
                  next_inspection_due=datetime.now(timezone.utc)
                  + timedelta(days=365 * jaren_vooruit))
        db.add(a)
        db.commit()
        db.refresh(a)
        return a.id
    finally:
        db.close()


def _melding(org_id, user_id, asset_id, categorie="schade"):
    db = SessionLocal()
    try:
        db.add(Melding(title="Scheur", organization_id=org_id, status="open",
                       priority="hoog", category=categorie, asset_id=asset_id,
                       created_by=user_id))
        db.commit()
    finally:
        db.close()


def _vul_mjop(user, project_id=None):
    """Een paar assets over meerdere jaren, met meldingen."""
    org = user.organization_id
    brug = _asset(org, user.id, code="BRG-001", asset_type="brug", score=5,
                  jaren_vooruit=1, project_id=project_id)
    _asset(org, user.id, code="VIA-010", asset_type="viaduct", score=6,
           jaren_vooruit=0, project_id=project_id)
    _asset(org, user.id, code="RIO-201", asset_type="riolering", score=4,
           jaren_vooruit=3, project_id=project_id, lengte=250)
    _asset(org, user.id, code="LMP-301", asset_type="verlichting", score=3,
           jaren_vooruit=5, project_id=project_id)
    _asset(org, user.id, code="PUT-901", asset_type="put", score=None,
           project_id=project_id)
    _melding(org, user.id, brug)
    _melding(org, user.id, brug, categorie="constructief")


def _xlsx(client, user, query="years=10"):
    r = client.get("/api/mjop/export.xlsx?" + query, headers=auth(user))
    assert r.status_code == 200, r.text
    return r, load_workbook(io.BytesIO(r.content))


def _kopregel(ws):
    return [c.value for c in ws[KOP_RIJ] if c.value is not None]


def _alle_tekst(wb):
    return " ".join(str(c.value) for ws in wb.worksheets for rij in ws.iter_rows()
                    for c in rij if c.value is not None)


def _pdf_inhoud(pdf: bytes) -> str:
    """De tekst uit de (gecomprimeerde) contentstreams van een fpdf2-PDF."""
    delen = []
    for stroom in re.findall(rb"stream\r?\n(.*?)\r?\nendstream", pdf, re.S):
        try:
            delen.append(zlib.decompress(stroom).decode("latin-1"))
        except zlib.error:
            continue
    return "".join(delen).replace("\\(", "(").replace("\\)", ")")


@pytest.fixture
def andere_gebruiker():
    """Admin van een tweede, losse organisatie."""
    db = SessionLocal()
    try:
        o = Organization(name="Andere Gemeente", plan=SubscriptionPlan.PROFESSIONAL,
                         status=AccountStatus.ACTIVE, max_users=10)
        db.add(o)
        db.commit()
        db.refresh(o)
        return _make_user(db, "admin@anderegemeente.nl", org=o,
                          role=UserRole.ADMIN, is_org_admin=True)
    finally:
        db.close()


# ─────────────────────────────────────────────────────────────────────────────
# Excel
# ─────────────────────────────────────────────────────────────────────────────

def test_xlsx_is_een_werkboek_met_briefhoofd_en_tabbladen(client, admin_user):
    _org_gegevens(admin_user.organization_id, name="Van Beek Infra B.V.",
                  kvk_number="24123456", contact_email="planning@vanbeek.nl")
    _vul_mjop(admin_user)

    r, wb = _xlsx(client, admin_user)
    assert r.headers["content-type"] == XLSX_MIME
    naam = re.search(r'filename="([^"]+)"', r.headers["content-disposition"]).group(1)
    assert naam.startswith("mjop-all-") and naam.endswith(".xlsx")

    assert wb.sheetnames == BLADEN
    for ws in wb.worksheets:
        briefhoofd = [c.value for rij in ws.iter_rows(min_row=1, max_row=4) for c in rij]
        assert "Van Beek Infra B.V." in briefhoofd, ws.title
        assert any("KvK 24123456" in str(v) for v in briefhoofd if v), ws.title
        assert str(ws["A7"].value).startswith("Meerjaren Onderhoudsplan"), ws.title


def test_xlsx_kopregel_is_die_van_de_pdf(client, admin_user):
    """Rij 11 van elk tabblad = de kolommen van dezelfde tabel in de PDF."""
    _vul_mjop(admin_user)
    _, wb = _xlsx(client, admin_user)
    k = _kolommen(met_index=False)
    for blad, sleutel in (("Kosten per jaar", "per_jaar"), ("MJOP-tabel", "tabel"),
                          ("Assets per type", "types"), ("Conditie-verdeling", "conditie"),
                          ("Meldingen per categorie", "categorieen"),
                          ("Assets met meldingen", "meldingen_assets")):
        assert _kopregel(wb[blad]) == [x.naam for x in k[sleutel]], blad

    # En die kolomkoppen staan ook echt in de PDF.
    pdf = client.get("/api/mjop/export.pdf?years=10", headers=auth(admin_user)).content
    tekst = _pdf_inhoud(pdf)
    for kolom in k["tabel"] + k["per_jaar"]:
        for woord in kolom.naam.split():
            assert woord.encode("cp1252").decode("latin-1") in tekst, kolom.naam


def test_xlsx_mjop_tabel_bevat_de_regels_en_een_totaal(client, admin_user):
    _vul_mjop(admin_user)
    preview = client.get("/api/mjop/preview?years=10", headers=auth(admin_user)).json()
    _, wb = _xlsx(client, admin_user)
    ws = wb["MJOP-tabel"]
    kop = _kopregel(ws)
    codes = [ws.cell(row=r, column=kop.index("Asset-code") + 1).value
             for r in range(KOP_RIJ + 1, KOP_RIJ + 1 + preview["count"])]
    assert codes == [x["asset_code"] for x in preview["items"]]
    totaalrij = KOP_RIJ + 1 + preview["count"]
    assert ws.cell(row=totaalrij, column=1).value == "Totaal"
    kol_max = kop.index(f"Max totaal (prijspeil {mjop.PRIJSPEIL_JAAR})") + 1
    assert ws.cell(row=totaalrij, column=kol_max).value == sum(
        x["max_total"] for x in preview["items"])


def test_xlsx_bedragen_zijn_getallen_met_euro_opmaak(client, admin_user):
    _vul_mjop(admin_user)
    _, wb = _xlsx(client, admin_user)
    for blad in ("MJOP-tabel", "Kosten per jaar"):
        ws = wb[blad]
        kop = _kopregel(ws)
        geld = [i + 1 for i, naam in enumerate(kop) if "totaal" in naam.lower()
                or "per eenheid" in naam]
        assert geld, blad
        for kolom in geld:
            c = ws.cell(row=KOP_RIJ + 1, column=kolom)
            assert isinstance(c.value, (int, float)) and not isinstance(c.value, bool), (blad, c.value)
            assert "€" in c.number_format, (blad, c.number_format)
        jaar = ws.cell(row=KOP_RIJ + 1, column=kop.index("Jaar") + 1)
        assert isinstance(jaar.value, int) and 2000 < jaar.value < 2100, jaar.value
        assert jaar.number_format == "0", "een jaartal hoort geen duizendtalpunt te krijgen"


def test_geindexeerde_kolommen_alleen_met_ingesteld_percentage(client, admin_user):
    _vul_mjop(admin_user)
    _, zonder = _xlsx(client, admin_user)
    assert not any("geïndexeerd" in k for k in _kopregel(zonder["MJOP-tabel"]))

    _org_gegevens(admin_user.organization_id, mjop_index_pct=3.0)
    _, met = _xlsx(client, admin_user)
    kop = _kopregel(met["MJOP-tabel"])
    assert kop[-2:] == ["Min geïndexeerd", "Max geïndexeerd"]
    assert _kopregel(met["Kosten per jaar"])[-2:] == ["Min geïndexeerd", "Max geïndexeerd"]
    toelichting = _alle_tekst(met)
    assert "3.0%" in toelichting, "de toelichting moet het gebruikte percentage noemen"

    # De PDF met zestien kolommen past nog steeds op de pagina.
    r = client.get("/api/mjop/export.pdf?years=10", headers=auth(admin_user))
    assert r.status_code == 200 and r.content[:4] == b"%PDF"


def test_xlsx_zonder_regels_is_een_leeg_maar_net_werkboek(client, admin_user):
    _, wb = _xlsx(client, admin_user)
    assert wb.sheetnames == BLADEN
    assert wb["Kosten per jaar"].cell(row=KOP_RIJ + 1, column=1).value == \
        "Geen gegevens voor deze selectie."


# ─────────────────────────────────────────────────────────────────────────────
# PDF
# ─────────────────────────────────────────────────────────────────────────────

def test_pdf_blijft_een_pdf_met_dezelfde_bestandsnaam(client, admin_user):
    pid = _project(admin_user.organization_id, admin_user.id, "Kunstwerken Noord")
    _vul_mjop(admin_user, project_id=pid)
    r = client.get(f"/api/mjop/export.pdf?years=10&project_id={pid}", headers=auth(admin_user))
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/pdf")
    assert r.content[:4] == b"%PDF"
    naam = re.search(r'filename="([^"]+)"', r.headers["content-disposition"]).group(1)
    assert re.fullmatch(r"mjop-kunstwerken-noord-\d{4}-\d{2}-\d{2}\.pdf", naam), naam
    tekst = _pdf_inhoud(r.content)
    for verwacht in ("Managementsamenvatting", "Kosten per jaar", "Volledige MJOP-tabel",
                     "Voorbeeld-berekening", "Bronnen en disclaimers", "BRG-001"):
        assert verwacht in tekst, verwacht


def test_pdf_zonder_regels(client, admin_user):
    r = client.get("/api/mjop/export.pdf?years=5", headers=auth(admin_user))
    assert r.status_code == 200, r.text
    assert r.content[:4] == b"%PDF"


# ─────────────────────────────────────────────────────────────────────────────
# Klantdata met logo en lastige tekens
# ─────────────────────────────────────────────────────────────────────────────

def test_logo_en_rare_tekens_breken_niets(client, admin_user):
    naam = "Gemeente Één → Beheer 🚧 € B.V."
    _org_gegevens(admin_user.organization_id, name=naam, logo_data_url=_logo_data_url(),
                  billing_address="Straat 1 → achterom\n1234 AB Plaats 🏗",
                  btw_number="NL001234567B01", contact_phone="+31 6 – 1234 5678")
    pid = _project(admin_user.organization_id, admin_user.id, "Wegen → 2026 🚧 € “zuid”")
    _asset(admin_user.organization_id, admin_user.id, code="BRG→01 🚧",
           asset_type="brug", score=4, project_id=pid, naam="Brug “De Één” – €")

    q = f"years=10&project_id={pid}"
    r = client.get("/api/mjop/export.pdf?" + q, headers=auth(admin_user))
    assert r.status_code == 200, r.text
    assert r.content[:4] == b"%PDF"
    r.headers["content-disposition"].encode("latin-1")   # header blijft geldig

    r, wb = _xlsx(client, admin_user, q)
    assert r.headers["content-disposition"].isascii()
    ws = wb["MJOP-tabel"]
    assert ws._images, "het logo hoort in het briefhoofd te staan"
    briefhoofd = [c.value for rij in ws.iter_rows(min_row=1, max_row=4) for c in rij]
    assert naam in briefhoofd
    assert "BRG→01 🚧" in _alle_tekst(wb)


def test_kapot_logo_breekt_niets(client, admin_user):
    _org_gegevens(admin_user.organization_id, logo_data_url="data:image/png;base64,bm9n")
    _vul_mjop(admin_user)
    assert client.get("/api/mjop/export.pdf", headers=auth(admin_user)).status_code == 200
    _xlsx(client, admin_user)


# ─────────────────────────────────────────────────────────────────────────────
# Rechten
# ─────────────────────────────────────────────────────────────────────────────

def test_andere_organisatie_ziet_niets(client, admin_user, andere_gebruiker):
    _org_gegevens(admin_user.organization_id, name="Eigenaar Infra")
    pid = _project(admin_user.organization_id, admin_user.id, "Geheim project")
    _asset(admin_user.organization_id, admin_user.id, code="GEHEIM-001",
           project_id=pid, naam="Geheime brug")

    for q in ("years=10", f"years=10&project_id={pid}"):
        r, wb = _xlsx(client, andere_gebruiker, q)
        alles = _alle_tekst(wb)
        assert "GEHEIM" not in alles and "Geheim" not in alles
        assert "Eigenaar Infra" not in alles
        assert "Andere Gemeente" in alles
        assert "geheim" not in r.headers["content-disposition"].lower()

        r = client.get("/api/mjop/export.pdf?" + q, headers=auth(andere_gebruiker))
        assert r.status_code == 200
        tekst = _pdf_inhoud(r.content)
        assert "GEHEIM" not in tekst and "Geheim" not in tekst
        assert "Eigenaar Infra" not in tekst


def test_zonder_login_geen_export(client):
    for pad in ("/api/mjop/export.xlsx", "/api/mjop/export.pdf"):
        assert client.get(pad).status_code in (401, 403), pad


def test_xlsx_valt_onder_de_kunstwerken_module(client, admin_user):
    _org_gegevens(admin_user.organization_id, enabled_modules=json.dumps([]))
    r = client.get("/api/mjop/export.xlsx", headers=auth(admin_user))
    assert r.status_code == 403
