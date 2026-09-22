"""Excel en PDF in de huisstijl voor tabellen uit het portaal, het dagboek en de auditlog.

Wat hier vastligt:

1. **Een tabel uit het portaal krijgt hetzelfde briefhoofd als de rapporten.**
   Logo en gegevens komen van de eigen organisatie, nooit uit het verzoek.
2. **Excel en PDF hebben dezelfde kolommen.**
3. **Het dagboek rekent in Nederlandse tijd**, met een totaal en een overzicht
   per project.
4. **De auditlog blijft alleen voor beheerders**, en elke export staat zelf
   weer in de auditlog.
5. **Te grote of scheve tabellen worden geweigerd met uitleg**, niet half
   gemaakt.
"""

import io

from openpyxl import load_workbook

from database import SessionLocal
from export_huisstijl import XLSX_MIME
from models import AuditLog, Organization

from .conftest import auth

TABEL = {
    "titel": "Meldingen rapportage",
    "ondertitel": "Project Noord",
    "bladen": [{
        "naam": "Meldingen",
        "kolommen": [{"naam": "Titel"}, {"naam": "Datum", "soort": "datum"},
                     {"naam": "Raming", "soort": "geld"}],
        "rijen": [["Gat in asfalt", "2026-09-12", 1234.5], ["=1+1", None, "88"]],
        "totaal": ["Totaal", None, 1322.5],
    }],
}


def _org_gegevens(user, **kw):
    db = SessionLocal()
    try:
        o = db.get(Organization, user.organization_id)
        o.billing_address = kw.get("adres", "Industrieweg 12\n3044 AS Rotterdam")
        o.kvk_number = kw.get("kvk", "24123456")
        db.commit()
        return o.name
    finally:
        db.close()


def _werkboek(r):
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == XLSX_MIME
    assert "attachment" in r.headers["content-disposition"]
    return load_workbook(io.BytesIO(r.content))


# ---------------------------------------------------------------------------
# Tabel uit het portaal
# ---------------------------------------------------------------------------

def test_tabel_als_excel_met_eigen_briefhoofd(client, admin_user):
    naam = _org_gegevens(admin_user)
    ws = _werkboek(client.post("/api/export/tabel.xlsx", headers=auth(admin_user), json=TABEL)).active
    assert ws["A1"].value == naam                        # geen logo: naam links
    assert ws["C2"].value == "Industrieweg 12, 3044 AS Rotterdam"
    assert ws["A7"].value == "Meldingen rapportage"
    assert [c.value for c in ws[11]] == ["Titel", "Datum", "Raming"]
    assert ws["C12"].value == 1234.5 and ws["C13"].value == 88.0
    assert ws["A13"].data_type == "s"                    # "=1+1" blijft tekst
    assert ws["A14"].value == "Totaal"


def test_tabel_als_pdf(client, admin_user):
    r = client.post("/api/export/tabel.pdf", headers=auth(admin_user), json=TABEL)
    assert r.status_code == 200, r.text
    assert r.content[:4] == b"%PDF" and r.headers["content-type"] == "application/pdf"


def test_briefhoofd_komt_niet_uit_het_verzoek(client, admin_user):
    naam = _org_gegevens(admin_user)
    body = {**TABEL, "klant": {"naam": "Iemand Anders B.V."}, "organization_id": "x"}
    ws = _werkboek(client.post("/api/export/tabel.xlsx", headers=auth(admin_user), json=body)).active
    assert ws["A1"].value == naam


def test_scheve_tabel_wordt_geweigerd(client, admin_user):
    scheef = {**TABEL, "bladen": [{**TABEL["bladen"][0], "rijen": [["alleen titel"]]}]}
    r = client.post("/api/export/tabel.xlsx", headers=auth(admin_user), json=scheef)
    assert r.status_code == 400 and "3 waarden" in r.json()["detail"]
    onbekend = {**TABEL, "bladen": [{**TABEL["bladen"][0],
                                     "kolommen": [{"naam": "A", "soort": "formule"}] * 3}]}
    assert client.post("/api/export/tabel.xlsx", headers=auth(admin_user),
                       json=onbekend).status_code == 400


def test_te_grote_pdf_wordt_geweigerd_met_uitleg(client, admin_user):
    groot = {**TABEL, "bladen": [{**TABEL["bladen"][0], "totaal": None,
                                  "rijen": [["x", None, 1]] * 3001}]}
    r = client.post("/api/export/tabel.pdf", headers=auth(admin_user), json=groot)
    assert r.status_code == 400 and "Verklein de selectie" in r.json()["detail"]


def test_tabel_export_vraagt_inlog(client):
    assert client.post("/api/export/tabel.xlsx", json=TABEL).status_code == 401


# ---------------------------------------------------------------------------
# Dagboek
# ---------------------------------------------------------------------------

def _dagboek(client, user):
    for titel, wanneer, minuten in (("Hotbox Stadhouderskade", "2026-09-14T06:30:00+00:00", 240),
                                    ("Scheuren vullen Ring", "2026-09-14T12:00:00+00:00", 90)):
        r = client.post("/api/daybook/entries", headers=auth(user), json={
            "title": titel, "occurred_at": wanneer, "duration_minutes": minuten,
            "entry_type": "manual_note"})
        assert r.status_code in (200, 201), r.text


def test_dagboek_als_excel_in_nederlandse_tijd(client, admin_user):
    _dagboek(client, admin_user)
    wb = _werkboek(client.get("/api/daybook/export.xlsx?from=2026-09-14&to=2026-09-14",
                              headers=auth(admin_user)))
    assert wb.sheetnames == ["Uren", "Per project"]
    ws = wb["Uren"]
    assert [c.value for c in ws[11]] == ["Datum", "Tijd", "Type", "Bron", "Activiteit",
                                         "Omschrijving", "Project", "Duur (min)", "Duur (uur)"]
    assert ws["B12"].value == "08:30"                     # 06:30 UTC = 08:30 zomertijd
    assert ws["C12"].value == "Notitie" and ws["E12"].value == "Hotbox Stadhouderskade"
    assert ws["E14"].value == "Totaal" and ws["H14"].value == 330 and ws["I14"].value == 5.5
    per_project = wb["Per project"]
    assert per_project["A12"].value == "(geen project)" and per_project["B12"].value == 330


def test_dagboek_als_pdf(client, admin_user):
    _dagboek(client, admin_user)
    r = client.get("/api/daybook/export.pdf?from=2026-09-14&to=2026-09-14", headers=auth(admin_user))
    assert r.status_code == 200 and r.content[:4] == b"%PDF"


def test_dagboek_van_een_ander_blijft_dicht(client, admin_user, viewer_user):
    r = client.get(f"/api/daybook/export.xlsx?from=2026-09-14&to=2026-09-14&user_id={admin_user.id}",
                   headers=auth(viewer_user))
    assert r.status_code == 403


# ---------------------------------------------------------------------------
# Auditlog
# ---------------------------------------------------------------------------

def test_auditlog_als_excel_en_pdf_voor_beheerder(client, admin_user):
    wb = _werkboek(client.get("/api/audit/logs/export.xlsx", headers=auth(admin_user)))
    ws = wb.active
    assert [c.value for c in ws[11]] == ["Tijd", "Actie", "Onderdeel", "ID", "Gebruiker",
                                         "IP-adres", "Details"]
    assert ws.page_setup.orientation == "landscape"
    r = client.get("/api/audit/logs/export.pdf", headers=auth(admin_user))
    assert r.status_code == 200 and r.content[:4] == b"%PDF"


def test_auditlog_export_alleen_voor_beheerders(client, viewer_user):
    for pad in ("/api/audit/logs/export.xlsx", "/api/audit/logs/export.pdf"):
        assert client.get(pad, headers=auth(viewer_user)).status_code == 403


def test_auditlog_export_staat_zelf_in_de_auditlog(client, admin_user):
    client.get("/api/audit/logs/export.xlsx", headers=auth(admin_user))
    db = SessionLocal()
    try:
        regels = db.query(AuditLog).filter(AuditLog.action == "audit.export.csv").all()
        assert any('"formaat": "xlsx"' in (r.details or "") for r in regels)
    finally:
        db.close()
