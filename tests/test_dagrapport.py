"""Projectdagrapport: dagboek, personeel, materieel, materiaal, afwijkingen, week.

Wat hier vastligt:

1. **Het dagrapport is van het project.** Iedereen in de organisatie leest het,
   iedereen behalve een lezer vult het aan. Een andere organisatie ziet niets.
2. **Materieel is dezelfde regel als in het werkdagboek.** Een machine die een
   medewerker op het project zet, staat in het dagrapport van iedereen, met
   dezelfde CO2 -- er wordt niets dubbel geteld.
3. **Standaardlijst en reductie.** Een machine uit de standaardlijst met alleen
   draaiuren levert een geschatte CO2 op; HVO100 levert een reductie ten
   opzichte van diesel op, uitgerekend over dezelfde liters.
4. **Het weekrapport** zet alles met een kolom per dag naast elkaar, en is als
   PDF en Excel te downloaden.
"""

import io
import json
from datetime import date, timedelta
from pathlib import Path

import pytest

import materieel as mt
from database import SessionLocal
from models import Organization, Project, User, UserRole
from tests.conftest import _make_user, auth

MAANDAG = date(2026, 9, 14)          # week 38 van 2026
DINSDAG = MAANDAG + timedelta(days=1)


def _project(org_id, maker_id, naam="Dagelijks onderhoud 2026"):
    db = SessionLocal()
    try:
        p = Project(name=naam, organization_id=org_id, created_by=maker_id)
        db.add(p)
        db.commit()
        return p.id
    finally:
        db.close()


@pytest.fixture
def project(org, admin_user):
    return _project(org.id, admin_user.id)


@pytest.fixture
def ander(request):
    """Een beheerder van een andere organisatie, met een eigen project."""
    db = SessionLocal()
    try:
        o = Organization(name=f"Ander-{request.node.name[:20]}")
        db.add(o)
        db.commit()
        u = _make_user(db, f"ander-{o.id[:8]}@test.nl", org=o, role=UserRole.ADMIN, is_org_admin=True)
        return u
    finally:
        db.close()


def _dag(client, user, project_id, datum=MAANDAG):
    r = client.get(f"/api/dagrapport/dag?project_id={project_id}&datum={datum}", headers=auth(user))
    assert r.status_code == 200, r.text
    return r.json()


# ---------------------------------------------------------------------------
# Dagboek
# ---------------------------------------------------------------------------

def test_log_weer_en_temperatuur_opslaan_en_teruglezen(client, admin_user, project):
    r = client.put("/api/dagrapport/dag", headers=auth(admin_user), json={
        "project_id": project, "datum": str(MAANDAG), "weer": "droog_zonnig",
        "temp_min": 15, "temp_max": 20,
        "log": "Locatie 1 begonnen met frezen om 07:00 uur.\nAfzetting geplaatst op 3 locaties."})
    assert r.status_code == 200, r.text
    d = _dag(client, admin_user, project)
    assert d["dag"]["weer_naam"] == "Droog, zonnig"
    assert d["dag"]["temp_min"] == 15 and d["dag"]["temp_max"] == 20
    assert "frezen om 07:00" in d["dag"]["log"]
    assert d["week"] == "2026-38"

    # Nog een keer opslaan werkt dezelfde dag bij, geen tweede dag.
    client.put("/api/dagrapport/dag", headers=auth(admin_user), json={
        "project_id": project, "datum": str(MAANDAG), "log": "Aangepast"})
    assert _dag(client, admin_user, project)["dag"]["log"] == "Aangepast"
    w = client.get(f"/api/dagrapport/week?project_id={project}&datum={MAANDAG}",
                   headers=auth(admin_user)).json()
    assert len(w["dagen"]) == 1


def test_onzin_in_het_dagboek_wordt_geweigerd(client, admin_user, project):
    basis = {"project_id": project, "datum": str(MAANDAG)}
    assert client.put("/api/dagrapport/dag", headers=auth(admin_user),
                      json={**basis, "weer": "hagelstorm"}).status_code == 400
    assert client.put("/api/dagrapport/dag", headers=auth(admin_user),
                      json={**basis, "temp_min": 25, "temp_max": 10}).status_code == 400
    assert client.put("/api/dagrapport/dag", headers=auth(admin_user),
                      json={**basis, "temp_max": 90}).status_code == 422


# ---------------------------------------------------------------------------
# Wie mag wat
# ---------------------------------------------------------------------------

def test_een_medewerker_vult_aan_en_een_lezer_niet(client, admin_user, technician_user,
                                                    viewer_user, project):
    r = client.put("/api/dagrapport/dag", headers=auth(technician_user), json={
        "project_id": project, "datum": str(MAANDAG), "log": "Bossages gesnoeid"})
    assert r.status_code == 200, r.text
    # De beheerder ziet wat de technicus schreef: het rapport is van het project.
    assert _dag(client, admin_user, project)["dag"]["log"] == "Bossages gesnoeid"
    assert _dag(client, admin_user, project)["dag"]["bijgewerkt_door"]

    r = client.put("/api/dagrapport/dag", headers=auth(viewer_user), json={
        "project_id": project, "datum": str(MAANDAG), "log": "mag niet"})
    assert r.status_code == 403
    assert client.post("/api/dagrapport/personeel", headers=auth(viewer_user), json={
        "project_id": project, "datum": str(MAANDAG), "naam": "X", "uren": 8}).status_code == 403
    # Lezen mag wel.
    assert _dag(client, viewer_user, project)["dag"]["log"] == "Bossages gesnoeid"
    cfg = client.get("/api/dagrapport/config", headers=auth(viewer_user)).json()
    assert cfg["mag_schrijven"] is False


def test_een_andere_organisatie_ziet_en_raakt_niets(client, admin_user, project, ander):
    p = client.post("/api/dagrapport/personeel", headers=auth(admin_user), json={
        "project_id": project, "datum": str(MAANDAG), "naam": "Piet", "uren": 8}).json()
    m = client.post("/api/dagrapport/materiaal", headers=auth(admin_user), json={
        "project_id": project, "datum": str(MAANDAG), "materiaal": "Asfalt", "hoeveelheid": 16}).json()
    a = client.post("/api/dagrapport/afwijkingen", headers=auth(admin_user), json={
        "project_id": project, "datum": str(MAANDAG), "omschrijving": "Meerwerk"}).json()

    h = auth(ander)
    assert client.get(f"/api/dagrapport/dag?project_id={project}&datum={MAANDAG}", headers=h).status_code == 404
    assert client.get(f"/api/dagrapport/week?project_id={project}&datum={MAANDAG}", headers=h).status_code == 404
    assert client.get(f"/api/dagrapport/week.pdf?project_id={project}&datum={MAANDAG}",
                      headers=h).status_code == 404
    assert client.put("/api/dagrapport/dag", headers=h, json={
        "project_id": project, "datum": str(MAANDAG), "log": "x"}).status_code == 404
    for pad, rid in (("personeel", p["id"]), ("materiaal", m["id"]), ("afwijkingen", a["id"])):
        assert client.patch(f"/api/dagrapport/{pad}/{rid}", headers=h, json={}).status_code == 404
        assert client.delete(f"/api/dagrapport/{pad}/{rid}", headers=h).status_code == 404
    # En de medewerkerslijst toont alleen de eigen organisatie.
    namen = [x["id"] for x in client.get("/api/dagrapport/config", headers=h).json()["medewerkers"]]
    assert admin_user.id not in namen


def test_verwijderen_mag_de_maker_of_een_beheerder(client, admin_user, technician_user,
                                                    inspector_user, project):
    r = client.post("/api/dagrapport/materiaal", headers=auth(technician_user), json={
        "project_id": project, "datum": str(MAANDAG), "materiaal": "Brekerzand", "hoeveelheid": 1,
        "eenheid": "stuk"}).json()
    assert client.delete(f"/api/dagrapport/materiaal/{r['id']}",
                         headers=auth(inspector_user)).status_code == 403
    assert client.delete(f"/api/dagrapport/materiaal/{r['id']}",
                         headers=auth(technician_user)).status_code == 200
    r = client.post("/api/dagrapport/materiaal", headers=auth(technician_user), json={
        "project_id": project, "datum": str(MAANDAG), "materiaal": "Brekerzand", "hoeveelheid": 1,
        "eenheid": "stuk"}).json()
    assert client.delete(f"/api/dagrapport/materiaal/{r['id']}",
                         headers=auth(admin_user)).status_code == 200


# ---------------------------------------------------------------------------
# Personeel
# ---------------------------------------------------------------------------

def test_eigen_medewerker_en_inhuur(client, admin_user, technician_user, project):
    r = client.post("/api/dagrapport/personeel", headers=auth(admin_user), json={
        "project_id": project, "datum": str(MAANDAG), "user_id": technician_user.id,
        "functie": "Machinist", "uren": 8})
    assert r.status_code == 201, r.text
    eigen = r.json()
    assert eigen["eigen"] and eigen["naam_toon"]            # naam uit het account
    assert eigen["bedrijf"]                                 # eigen organisatie

    r = client.post("/api/dagrapport/personeel", headers=auth(admin_user), json={
        "project_id": project, "datum": str(MAANDAG), "functie": "Verkeersregelaar",
        "bedrijf": "Buko Infrasupport B.V.", "uren": 9})
    assert r.status_code == 201, r.text                     # naam van inhuur is vaak onbekend

    assert client.post("/api/dagrapport/personeel", headers=auth(admin_user), json={
        "project_id": project, "datum": str(MAANDAG), "uren": 8}).status_code == 400
    assert client.post("/api/dagrapport/personeel", headers=auth(admin_user), json={
        "project_id": project, "datum": str(MAANDAG), "naam": "X", "uren": 25}).status_code == 422
    assert client.post("/api/dagrapport/personeel", headers=auth(admin_user), json={
        "project_id": project, "datum": str(MAANDAG), "naam": "X", "uren": -1}).status_code == 422

    d = _dag(client, admin_user, project)
    assert d["uren_personeel"] == 17


def test_ploeg_van_gisteren_overnemen_zonder_dubbelen(client, admin_user, technician_user, project):
    for body in ({"user_id": technician_user.id, "functie": "Machinist", "uren": 8},
                 {"naam": "Armando", "functie": "Vakman GWW", "bedrijf": "Infrakracht", "uren": 8}):
        client.post("/api/dagrapport/personeel", headers=auth(admin_user),
                    json={"project_id": project, "datum": str(MAANDAG), **body})
    kopie = {"project_id": project, "van": str(MAANDAG), "naar": str(DINSDAG)}
    r = client.post("/api/dagrapport/personeel/kopieer", headers=auth(admin_user), json=kopie)
    assert r.status_code == 200, r.text
    assert r.json()["toegevoegd"] == 2
    # Twee keer tikken levert geen dubbele uren op.
    r = client.post("/api/dagrapport/personeel/kopieer", headers=auth(admin_user), json=kopie)
    assert r.json()["toegevoegd"] == 0 and r.json()["overgeslagen"] == 2
    assert _dag(client, admin_user, project, DINSDAG)["uren_personeel"] == 16


# ---------------------------------------------------------------------------
# Materieel: standaardlijst, CO2 en reductie
# ---------------------------------------------------------------------------

def test_standaardlijst_staat_in_de_config_met_verbruik_en_vermogen(client, admin_user):
    cfg = client.get("/api/materieel/config", headers=auth(admin_user)).json()
    lijst = {m["code"]: m for m in cfg["catalogus"]["materieel"]}
    assert len(lijst) >= 30
    kraan = lijst["mobiele_kraan_15t"]
    assert kraan["vermogen_kw"] == 105
    assert kraan["verbruik_per_uur"] == pytest.approx(105 * kraan["belasting"] * 0.24, abs=0.06)
    for m in lijst.values():
        assert m["soort"] in mt.SOORTEN, m
        assert m["energiedrager"] in mt.ENERGIEDRAGERS, m
        assert m["verbruik_per_uur"] > 0


def test_machine_uit_de_standaardlijst_geeft_een_geschatte_co2(client, admin_user, project):
    r = client.post("/api/materieel/inzet", headers=auth(admin_user), json={
        "catalogus_code": "mobiele_kraan_15t", "project_id": project, "datum": str(MAANDAG),
        "draaiuren": 9, "energiedrager": "hvo100", "emissieklasse": "stage_v",
        "leverancier": "Niek Konijn BV"})
    assert r.status_code == 201, r.text
    regel = r.json()
    assert regel["materieel_naam"].startswith("Mobiele kraan")
    assert regel["vermogen_kw"] == 105 and regel["emissieklasse_naam"] == "Stage V"
    assert regel["co2_methode"] == "geschat"
    liters = 9 * regel["verbruik_per_uur"]
    hvo = mt.factor_voor("hvo100").kg_co2_per_eenheid
    diesel = mt.factor_voor("diesel").kg_co2_per_eenheid
    assert regel["co2_kg"] == pytest.approx(liters * hvo, rel=1e-3)
    # De reductie is wat dezelfde liters op diesel hadden uitgestoten, min de uitstoot.
    assert regel["reductie_kg"] == pytest.approx(liters * (diesel - hvo), rel=1e-3)

    assert client.post("/api/materieel/inzet", headers=auth(admin_user), json={
        "catalogus_code": "bestaat_niet", "draaiuren": 1}).status_code == 400


def test_diesel_heeft_geen_reductie_en_stroom_geen_vergelijking(client, admin_user):
    assert mt.reductie_kg(energiedrager="diesel", co2_kg=32.5, co2_factor=3.251) == 0
    assert mt.reductie_kg(energiedrager="elektrisch", co2_kg=10, co2_factor=0.483) is None
    assert mt.reductie_kg(energiedrager="hvo100", co2_kg=None, co2_factor=0.468) is None


def test_eigen_verbruik_wint_van_de_standaardlijst(client, admin_user, project):
    r = client.post("/api/materieel/inzet", headers=auth(admin_user), json={
        "catalogus_code": "asfaltfrees_05", "project_id": project, "datum": str(MAANDAG),
        "draaiuren": 3, "verbruik_per_uur": 10})
    assert r.json()["verbruik_per_uur"] == 10
    # Getankte liters gaan voor alles: gemeten.
    r = client.post("/api/materieel/inzet", headers=auth(admin_user), json={
        "catalogus_code": "asfaltfrees_05", "project_id": project, "datum": str(MAANDAG),
        "draaiuren": 3, "brandstof_hoeveelheid": 40})
    assert r.json()["co2_methode"] == "gemeten"


def test_materieel_van_een_collega_staat_in_het_dagrapport(client, admin_user, technician_user, project):
    """In zijn eigen werkdagboek ziet ieder alleen zijn eigen machines; in het
    dagrapport van het project staan ze allemaal, met dezelfde CO2."""
    client.post("/api/materieel/inzet", headers=auth(technician_user), json={
        "catalogus_code": "kipper_4as", "project_id": project, "datum": str(MAANDAG),
        "brandstof_hoeveelheid": 50})
    client.post("/api/materieel/inzet", headers=auth(admin_user), json={
        "catalogus_code": "mobiele_kraan_15t", "project_id": project, "datum": str(MAANDAG),
        "brandstof_hoeveelheid": 30})
    d = _dag(client, admin_user, project)
    assert len(d["materieel"]) == 2
    diesel = mt.factor_voor("diesel").kg_co2_per_eenheid
    assert d["co2"]["uitstoot_kg"] == pytest.approx(80 * diesel, abs=0.01)
    assert d["co2"]["reductie_kg"] == 0


# ---------------------------------------------------------------------------
# Materiaal en afwijkingen
# ---------------------------------------------------------------------------

def test_materiaal_en_afwijking(client, admin_user, project):
    r = client.post("/api/dagrapport/materiaal", headers=auth(admin_user), json={
        "project_id": project, "datum": str(MAANDAG), "leverancier": "Bnext.nl Amsterdam B.V.",
        "materiaal": "TH Asfalt", "richting": "afvoer", "hoeveelheid": 16, "eenheid": "ton"})
    assert r.status_code == 201, r.text
    assert r.json()["richting_naam"] == "Afvoer"
    assert client.post("/api/dagrapport/materiaal", headers=auth(admin_user), json={
        "project_id": project, "datum": str(MAANDAG), "materiaal": "x", "hoeveelheid": 1,
        "eenheid": "emmers"}).status_code == 400
    assert client.post("/api/dagrapport/materiaal", headers=auth(admin_user), json={
        "project_id": project, "datum": str(MAANDAG), "materiaal": "  ", "hoeveelheid": 1}).status_code == 400

    r = client.post("/api/dagrapport/afwijkingen", headers=auth(admin_user), json={
        "project_id": project, "datum": str(MAANDAG),
        "omschrijving": "Overhangende bossages over het fietspad",
        "maatregel": "Eerste uur bossages snoeien met kettingzaag",
        "adres": "Baanakkerspad, 1024 PH Amsterdam", "stagnatie": "vertraging",
        "duur_uren": 1, "soort": "meerwerk", "bedrag": 500})
    assert r.status_code == 201, r.text
    afw = r.json()
    r = client.patch(f"/api/dagrapport/afwijkingen/{afw['id']}", headers=auth(admin_user),
                     json={"bedrag": 650})
    assert r.json()["bedrag"] == 650
    assert client.patch(f"/api/dagrapport/afwijkingen/{afw['id']}", headers=auth(admin_user),
                        json={"soort": "cadeau"}).status_code == 400


# ---------------------------------------------------------------------------
# Week
# ---------------------------------------------------------------------------

def _vul_week(client, admin_user, technician_user, project):
    client.put("/api/dagrapport/dag", headers=auth(admin_user), json={
        "project_id": project, "datum": str(MAANDAG), "weer": "droog_zonnig",
        "temp_min": 20, "temp_max": 25, "log": "Locaties vastgelegd"})
    client.put("/api/dagrapport/dag", headers=auth(admin_user), json={
        "project_id": project, "datum": str(DINSDAG), "weer": "droog_bewolkt",
        "temp_min": 15, "temp_max": 20, "log": "Frezen en asfalteren"})
    for datum, uren in ((MAANDAG, 3), (DINSDAG, 8)):
        client.post("/api/dagrapport/personeel", headers=auth(admin_user), json={
            "project_id": project, "datum": str(datum), "user_id": technician_user.id,
            "functie": "Uitvoerder", "uren": uren})
    client.post("/api/dagrapport/personeel", headers=auth(admin_user), json={
        "project_id": project, "datum": str(DINSDAG), "functie": "Verkeersregelaar",
        "bedrijf": "Buko", "uren": 9})
    client.post("/api/materieel/inzet", headers=auth(admin_user), json={
        "catalogus_code": "asfaltfrees_05", "project_id": project, "datum": str(DINSDAG),
        "draaiuren": 3, "energiedrager": "hvo100", "emissieklasse": "stage_v",
        "leverancier": "Aduco"})
    client.post("/api/materieel/inzet", headers=auth(admin_user), json={
        "catalogus_code": "mobiele_kraan_15t", "project_id": project, "datum": str(DINSDAG),
        "draaiuren": 9, "energiedrager": "hvo100", "emissieklasse": "stage_v"})
    client.post("/api/dagrapport/materiaal", headers=auth(admin_user), json={
        "project_id": project, "datum": str(DINSDAG), "leverancier": "Bnext",
        "materiaal": "TH Asfalt", "richting": "afvoer", "hoeveelheid": 16})
    client.post("/api/dagrapport/afwijkingen", headers=auth(admin_user), json={
        "project_id": project, "datum": str(MAANDAG), "omschrijving": "Voetpad meeasfalteren",
        "soort": "meerwerk", "bedrag": 2000, "stagnatie": "vertraging", "duur_uren": 2})
    # Buiten de week: telt niet mee.
    client.post("/api/dagrapport/personeel", headers=auth(admin_user), json={
        "project_id": project, "datum": str(MAANDAG - timedelta(days=1)), "naam": "Zondag ervoor",
        "uren": 8})


def test_de_week_met_een_kolom_per_dag(client, admin_user, technician_user, project):
    _vul_week(client, admin_user, technician_user, project)
    w = client.get(f"/api/dagrapport/week?project_id={project}&datum={DINSDAG + timedelta(days=3)}",
                   headers=auth(admin_user)).json()
    assert w["week"] == "2026-38" and w["maandag"] == str(MAANDAG)
    assert len(w["dagen"]) == 2

    uitvoerder = next(g for g in w["personeel"] if g["voorbeeld"]["functie"] == "Uitvoerder")
    assert uitvoerder["dagen"][:2] == [3, 8] and uitvoerder["totaal"] == 11
    assert w["uren_personeel"] == 20                     # 3 + 8 + 9, niet de zondag ervoor

    assert len(w["materieel"]) == 2 and w["uren_materieel"] == 12
    assert w["co2"]["uitstoot_kg"] > 0 and w["co2"]["reductie_kg"] > w["co2"]["uitstoot_kg"]
    assert w["co2"]["waarvan_geschat_kg"] == w["co2"]["uitstoot_kg"]
    assert w["materiaal"][0]["dagen"][1] == 16
    assert w["meerwerk_bedrag"] == 2000


def test_weekrapport_als_pdf_en_excel(client, admin_user, technician_user, project):
    _vul_week(client, admin_user, technician_user, project)
    r = client.get(f"/api/dagrapport/week.pdf?project_id={project}&datum={MAANDAG}",
                   headers=auth(admin_user))
    assert r.status_code == 200, r.text
    assert r.content.startswith(b"%PDF")
    assert "Weekrapport" in r.headers["content-disposition"] and "2026-38" in r.headers["content-disposition"]

    r = client.get(f"/api/dagrapport/week.xlsx?project_id={project}&datum={MAANDAG}",
                   headers=auth(admin_user))
    assert r.status_code == 200, r.text
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.content))
    assert {"Dagboek", "Personeel", "Materieel", "Materiaal", "Afwijkingen"} <= set(wb.sheetnames)
    tekst = " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value)
    assert "Frezen en asfalteren" in tekst and "Verkeersregelaar" in tekst and "TH Asfalt" in tekst


def test_lege_week_geeft_toch_een_rapport(client, admin_user, project):
    r = client.get(f"/api/dagrapport/week.pdf?project_id={project}&datum={MAANDAG}",
                   headers=auth(admin_user))
    assert r.status_code == 200 and r.content.startswith(b"%PDF")


def test_catalogus_bestand_is_geldig():
    data = json.loads(Path("data/materieel_catalogus.json").read_text(encoding="utf-8"))
    codes = [m["code"] for m in data["materieel"]]
    assert len(codes) == len(set(codes))
    assert data["toelichting"]


def test_project_verwijderd_dan_is_het_dagrapport_onvindbaar(client, admin_user, project):
    client.put("/api/dagrapport/dag", headers=auth(admin_user), json={
        "project_id": project, "datum": str(MAANDAG), "log": "x"})
    assert client.delete(f"/api/projects/{project}?hard=true", headers=auth(admin_user)).status_code == 200
    assert client.get(f"/api/dagrapport/dag?project_id={project}&datum={MAANDAG}",
                      headers=auth(admin_user)).status_code == 404
    db = SessionLocal()
    try:
        assert db.query(User).count() >= 1                 # niets anders omgevallen
    finally:
        db.close()
