"""Eén huisstijl voor PDF en Excel: FieldOps-kleuren, logo en gegevens van de klant.

Wat hier vastligt:

1. **Het briefhoofd is van de klant.** Naam, adres, KvK, BTW en contact komen
   uit de organisatie; wat leeg is valt weg in plaats van een kaal "KvK ".
2. **Een kapot logo breekt geen export.** Dan staat de naam er.
3. **Excel leest als de PDF.** Zelfde briefhoofd, blauwe kopregel, gestreepte
   rijen, en de voet "Pagina x van n".
4. **Excel rekent.** Bedragen, getallen en datums zijn echte waarden met een
   opmaak, geen tekst.
5. **Een cel die met = begint is tekst, geen formule.** Wat een gebruiker in
   een titel typt, mag in Excel niets uitvoeren.
"""

import base64
import io
import re
from datetime import date, datetime

import pytest
from openpyxl import load_workbook
from PIL import Image

import export_huisstijl as h


class _Org:
    def __init__(self, **kw):
        self.name = kw.get("name", "Van Beek Infra B.V.")
        self.logo_data_url = kw.get("logo_data_url")
        self.billing_address = kw.get("billing_address", "Industrieweg 12\n3044 AS Rotterdam")
        self.kvk_number = kw.get("kvk_number", "24123456")
        self.btw_number = kw.get("btw_number", "")
        self.contact_email = kw.get("contact_email", "planning@vanbeek.nl")
        self.contact_phone = kw.get("contact_phone", None)


def _png_data_url(b=300, hoogte=100):
    buf = io.BytesIO()
    Image.new("RGB", (b, hoogte), (22, 101, 52)).save(buf, "PNG")
    return "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode()


KOLOMMEN = [h.Kolom("Titel"), h.Kolom("Gemeld op", "datum"), h.Kolom("Oppervlak m²", "getal", 1),
            h.Kolom("Raming", "geld"), h.Kolom("Aantal", "heel")]
RIJEN = [["Gat in asfalt", "2026-09-12T10:00:00", 12.5, 1234.5, 3],
         ["=HYPERLINK(\"http://x\")", date(2026, 9, 1), None, "88", 1]]


def _blad(**kw):
    return h.Blad("Meldingen", KOLOMMEN, RIJEN, **kw)


# ---------------------------------------------------------------------------
# De klant
# ---------------------------------------------------------------------------

def test_gegevens_van_de_klant_in_het_briefhoofd():
    k = h.klant_van(_Org())
    assert k.naam == "Van Beek Infra B.V."
    assert k.gegevens() == ["Industrieweg 12, 3044 AS Rotterdam", "KvK 24123456", "planning@vanbeek.nl"]


def test_lege_gegevens_vallen_weg():
    k = h.klant_van(_Org(billing_address=None, kvk_number=None, contact_email=None))
    assert k.gegevens() == []
    assert h.klant_van(None).naam == "FieldOps"


def test_logo_wordt_png():
    logo = h.klant_van(_Org(logo_data_url=_png_data_url())).logo
    assert logo.png.startswith(b"\x89PNG") and round(logo.verhouding) == 3


@pytest.mark.parametrize("kapot", ["data:image/png;base64,bm9n", "data:image/png;base64,@@@",
                                   "https://example.com/logo.png", ""])
def test_kapot_logo_breekt_niets(kapot):
    k = h.klant_van(_Org(logo_data_url=kapot))
    assert k.logo is None
    assert h.pdf_van(k, "Proef", [_blad()]).startswith(b"%PDF")
    assert h.excel_van(k, "Proef", [_blad()])[:2] == b"PK"


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def _pdf_tekst(pdf: h.HuisstijlPDF) -> str:
    pdf.compress = False
    return pdf.uitvoer().decode("latin-1")


def test_pdf_briefhoofd_titel_en_voet():
    k = h.klant_van(_Org())
    pdf = h.HuisstijlPDF(k, "Meldingen rapportage", ondertitel="Project Noord")
    pdf.add_page()
    pdf.titelblok()
    pdf.blad(_blad(totaal=["Totaal", "", 12.5, 1322.5, 4]))
    ruw = _pdf_tekst(pdf)
    for verwacht in ("Van Beek Infra B.V.", "KvK 24123456", "Meldingen rapportage",
                     "Pagina 1 van", "Opgesteld met FieldOps", "Gat in asfalt", "1.234,50"):
        assert verwacht in ruw, verwacht


def test_pdf_tekens_buiten_de_letter_worden_leesbaar():
    assert h.pdf_tekst("m² → € 5 – “ok” ✓ ≥ 3 🚧") == "m² -> € 5 – “ok” v >= 3 "
    pdf = h.HuisstijlPDF(h.klant_van(_Org()), "Proef → 🚧")
    pdf.add_page()
    pdf.tekst("Verzakking ≥ 5 cm ✓ ok")
    assert pdf.uitvoer().startswith(b"%PDF")


def test_pdf_lange_tabel_loopt_over_paginas_door():
    rijen = [["Melding %d" % i, "2026-09-01", i, i * 10, 1] for i in range(120)]
    uit = h.pdf_van(h.klant_van(_Org()), "Lang", [h.Blad("Alles", KOLOMMEN, rijen)])
    assert len(re.findall(rb"/Type /Page\b", uit)) >= 3


def test_waarden_in_nederlandse_notatie():
    assert h.als_tekst(h.Kolom("x", "geld"), 1234.5) == "€ 1.234,50"
    assert h.als_tekst(h.Kolom("x", "getal", 1), 1234.56) == "1.234,6"
    assert h.als_tekst(h.Kolom("x", "procent"), 0.913) == "91%"
    assert h.als_tekst(h.Kolom("x", "datum"), "2026-09-21T23:30:00+00:00") == "22-09-2026"
    assert h.als_tekst(h.Kolom("x"), None) == ""
    assert h.als_tekst(h.Kolom("x"), True) == "ja"


def test_rij_met_verkeerd_aantal_waarden_valt_op():
    with pytest.raises(ValueError):
        h.Blad("Fout", KOLOMMEN, [["alleen titel"]])


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

def _werkboek(**kw):
    k = h.klant_van(_Org(logo_data_url=kw.pop("logo", _png_data_url())))
    return load_workbook(io.BytesIO(h.excel_van(k, "Meldingen rapportage", [_blad(**kw)],
                                                ondertitel="Project Noord")))


def test_excel_briefhoofd_en_titel_als_de_pdf():
    ws = _werkboek().active
    assert ws["E1"].value == "Van Beek Infra B.V."            # rechts, naast het logo
    assert ws["E2"].value == "Industrieweg 12, 3044 AS Rotterdam"
    assert len(ws._images) == 1
    assert ws["A5"].fill.start_color.rgb.endswith(h.BLAUW)   # de blauwe lijn
    assert ws["A7"].value == "Meldingen rapportage"
    assert ws["A8"].value == "Project Noord"


def test_excel_zonder_logo_zet_de_naam_links():
    ws = _werkboek(logo=None).active
    assert ws["A1"].value == "Van Beek Infra B.V." and not ws._images


def test_excel_kopregel_blauw_en_bevroren():
    ws = _werkboek().active
    kop = [c.value for c in ws[11]]
    assert kop == [k.naam for k in KOLOMMEN]
    assert ws["A11"].fill.start_color.rgb.endswith(h.BLAUW)
    assert ws["A11"].font.color.rgb.endswith(h.WIT) and ws["A11"].font.b
    assert ws.freeze_panes == "A12"
    assert ws.auto_filter.ref == "A11:E13"
    assert ws.print_title_rows == "$11:$11"
    assert "Pagina &P van &N" in ws.oddFooter.right.text


def test_excel_rekent_met_echte_waarden():
    ws = _werkboek().active
    assert ws["B12"].value == datetime(2026, 9, 12) and ws["B12"].number_format == "dd-mm-yyyy"
    assert ws["D12"].value == 1234.5 and "€" in ws["D12"].number_format
    assert ws["D13"].value == 88.0                            # "88" als tekst wordt een getal
    assert ws["C13"].value is None


def test_excel_formule_blijft_tekst():
    ws = _werkboek().active
    assert ws["A13"].value.startswith("=HYPERLINK") and ws["A13"].data_type == "s"


def test_excel_totaal_en_toelichting():
    ws = _werkboek(totaal=["Totaal", "", 12.5, 1322.5, 4],
                   toelichting=["Prijspeil 2026, exclusief BTW."]).active
    assert ws["A14"].value == "Totaal" and ws["A14"].font.b
    assert ws["A16"].value == "Prijspeil 2026, exclusief BTW."


def test_excel_meerdere_tabbladen_met_geldige_namen():
    k = h.klant_van(_Org())
    uit = h.excel_van(k, "MJOP", [h.Blad("Per jaar / totaal", KOLOMMEN, RIJEN),
                                   h.Blad("Per jaar / totaal", KOLOMMEN, [])])
    wb = load_workbook(io.BytesIO(uit))
    assert wb.sheetnames == ["Per jaar - totaal", "Per jaar - totaal (2)"]
    assert wb.worksheets[1]["A12"].value == "Geen gegevens voor deze selectie."


def test_bestandsnaam_zonder_rare_tekens():
    naam = h.bestandsnaam("MJOP", "Gemeente Één / Noord", ext="xlsx")
    assert re.fullmatch(r"MJOP-Gemeente-Een-Noord-\d{4}-\d{2}-\d{2}\.xlsx", naam)


# ---------------------------------------------------------------------------
# Breedtes, afbreken en grote tabellen
# ---------------------------------------------------------------------------

def test_jaar_en_hele_euros():
    assert h.als_tekst(h.Kolom("Jaar", "jaar"), 2026) == "2026"
    assert h.als_tekst(h.Kolom("x", "geld", decimalen=0), 1234567.4) == "€ 1.234.567"
    assert h.Kolom("Jaar", "jaar").excel_formaat == "0"
    assert h.Kolom("x", "geld", decimalen=0).excel_formaat == '"€" #,##0'


def test_lange_woorden_breken_af_met_streepje():
    uit = h.afbreekbaar("Ontruimingsalarminstallatie bij de hoofdingang")
    assert uit.replace("\u00ad", "") == "Ontruimingsalarminstallatie bij de hoofdingang"
    assert "Ontrui\u00admings" in uit and "hoofdingang" in uit     # korte woorden blijven heel


def test_getallen_breken_nooit_af():
    pdf = h.HuisstijlPDF(h.klant_van(_Org()), "Breedtes")
    pdf.add_page()
    kol = [h.Kolom("Omschrijving"), h.Kolom("Bedrag", "geld")]
    breedtes, _ = h.kolombreedtes(pdf, kol, [["x " * 80, 1234567.89]], None, 8.5)
    pdf.set_font(h.LETTER_PDF, "", 8.5)
    assert breedtes[1] >= pdf.get_string_width("€ 1.234.567,89") + 3


def test_grote_tabel_is_snel_en_zegt_dat_hij_inkort():
    rijen = [[f"Melding {i} " + "met een hele lange omschrijving " * 6, "2026-09-01", i, i * 10, 1]
             for i in range(h.SNEL_VANAF + 50)]
    pdf = h.HuisstijlPDF(h.klant_van(_Org()), "Groot")
    pdf.add_page()
    pdf.tabel(KOLOMMEN, rijen)
    ruw = _pdf_tekst(pdf)
    assert "de Excel-export bevat ze volledig" in ruw
    assert "Pagina 1 van" in ruw


def test_excel_codes_breken_niet_af_en_tekst_krijgt_hoogte():
    k = h.klant_van(_Org())
    kol = [h.Kolom("Code"), h.Kolom("Toelichting"), h.Kolom("Bedrag", "geld")]
    rijen = [["BRUG.VOEGOVERGANGEN", "woord " * 60, 1234567.89], ["BRUG.LEUNING", "kort", 5]]
    ws = load_workbook(io.BytesIO(h.excel_van(k, "Proef", [h.Blad("Blad", kol, rijen)]))).active
    assert not ws["A12"].alignment.wrap_text
    assert ws.column_dimensions["A"].width >= len("BRUG.VOEGOVERGANGEN")
    assert ws.column_dimensions["C"].width >= len("€ 1.234.567,89")
    assert ws.row_dimensions[12].height > 24                 # afgebroken toelichting
    assert ws.row_dimensions[13].height is None              # één regel: Excel zelf


def test_excel_smal_blad_schuift_niet_over_het_logo_en_toelichting_loopt_terug():
    k = h.klant_van(_Org(logo_data_url=_png_data_url()))
    blad = h.Blad("Smal", [h.Kolom("A"), h.Kolom("B")], [["x", "y"]],
                  toelichting=["Een lange toelichting " * 12])
    ws = load_workbook(io.BytesIO(h.excel_van(k, "Proef", [blad]))).active
    totaal = sum(ws.column_dimensions[c].width for c in "AB")
    assert totaal >= 58
    assert "A14:B14" in [str(r) for r in ws.merged_cells.ranges]
    assert ws["A14"].alignment.wrap_text


@pytest.mark.parametrize("aantal", [120, h.SNEL_VANAF + 50])
def test_liggende_tabel_blijft_liggend_op_vervolgpaginas(aantal):
    # Een liggende tabel in een staand rapport (zoals de MJOP-tabel): elke
    # vervolgpagina moet ook liggend zijn, anders valt de tabel rechts af.
    rijen = [["Melding %d" % i, "2026-09-01", i, i * 10, 1] for i in range(aantal)]
    pdf = h.HuisstijlPDF(h.klant_van(_Org()), "Rapport")
    pdf.add_page()
    pdf.tekst("Staande inleiding")
    pdf.add_page(orientation="L")
    pdf.tabel(KOLOMMEN, rijen)
    maten = [(float(b), float(hg)) for b, hg in
             re.findall(rb"/MediaBox \[0 0 ([\d.]+) ([\d.]+)\]", pdf.uitvoer())]
    assert len(maten) >= 4 and maten[0][0] < maten[0][1]          # eerste pagina staand
    assert all(b > hg for b, hg in maten[1:])
