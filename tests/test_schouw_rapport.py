"""Schouwrapport als PDF: elke schade met het beeld erbij.

Wat hier vastligt:

1. **Elke schade die niet is afgewezen staat erin, met de foto** en het rode
   vak op de plek waar hij is gezien.
2. **Afgewezen schades staan achteraan in een lijst**, zonder foto.
3. **Een ronde van een andere organisatie is onvindbaar.**
"""

import io
import re
import zlib

import pytest

import schouw_leren
import schouw_rapport
from database import SessionLocal
from models import Schouwrit, Schouwwaarneming

from .conftest import auth
from .test_schouw_rijstand import KUIL, _jpeg, _opname, _rit, _schades, nep_model


@pytest.fixture
def model(monkeypatch):
    return nep_model(monkeypatch)


@pytest.fixture(autouse=True)
def _schone_voorbeelden():
    schouw_leren._cache.clear()
    yield
    schouw_leren._cache.clear()


def _pdf_tekst(inhoud: bytes) -> str:
    """De tekst uit de (door fpdf2 gecomprimeerde) inhoud van de pagina's."""
    delen = []
    for stuk in re.findall(rb"stream\r?\n(.*?)\r?\nendstream", inhoud, flags=re.S):
        try:
            delen.append(zlib.decompress(stuk).decode("latin-1"))
        except zlib.error:
            continue
    return "\n".join(delen)


def _aantal_beelden(inhoud: bytes) -> int:
    return inhoud.count(b"/Subtype /Image")


def test_rapport_heeft_elke_schade_met_beeld(client, admin_user, model):
    rit_id = _rit(client, admin_user)
    model["antwoorden"] += [
        {"bruikbaar": True, "wegschade": [dict(KUIL)]},
        {"bruikbaar": True, "wegschade": [dict(KUIL, schadebeeld="rafeling", ernst="E",
                                              kader=[0.1, 0.6, 0.3, 0.9])]},
        {"bruikbaar": True, "wegschade": [dict(KUIL, schadebeeld="spoorvorming", ernst="L")]},
    ]
    # Ver uit elkaar in tijd, zodat het drie schades blijven.
    _opname(client, admin_user, rit_id, 1, gemaakt_op="2026-09-22T10:00:00+00:00")
    _opname(client, admin_user, rit_id, 2, gemaakt_op="2026-09-22T10:05:00+00:00")
    _opname(client, admin_user, rit_id, 3, gemaakt_op="2026-09-22T10:10:00+00:00")
    schades = _schades(rit_id)
    assert len(schades) == 3
    afgewezen = next(w for w in schades if w.crow_schadebeeld == "spoorvorming")
    r = client.patch(f"/api/schouw/waarnemingen/{afgewezen.id}", headers=auth(admin_user),
                     json={"afgewezen": True, "afwijs_reden": "schaduw"})
    assert r.status_code == 200, r.text

    r = client.get(f"/api/schouw/ritten/{rit_id}/rapport.pdf", headers=auth(admin_user))
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == "application/pdf"
    assert "Schouwrapport-Haarlem-Zuid" in r.headers["content-disposition"]
    assert r.content.startswith(b"%PDF")

    tekst = _pdf_tekst(r.content)
    assert "Schouwrapport" in tekst
    assert "Kuil" in tekst and "Rafeling" in tekst
    assert "Afgewezen" in tekst and "Schaduw" in tekst
    # Twee schades die meetellen, elk met hun beeld; de afgewezen zonder.
    assert _aantal_beelden(r.content) == 2


def test_schade_zonder_bewaard_beeld_staat_er_toch_in(client, admin_user):
    rit_id = client.post("/api/schouw/ritten", headers=auth(admin_user),
                         json={"gebied": "Centrum"}).json()["id"]
    db = SessionLocal()
    try:
        rit = db.get(Schouwrit, rit_id)
        db.add(Schouwwaarneming(schouwrit_id=rit.id, organization_id=rit.organization_id,
                                crow_verharding="asfalt", crow_schadebeeld="kuilen",
                                crow_ernst="E", bron="handmatig", bevestigd=True))
        db.commit()
    finally:
        db.close()
    r = client.get(f"/api/schouw/ritten/{rit_id}/rapport.pdf", headers=auth(admin_user))
    assert r.status_code == 200, r.text
    tekst = _pdf_tekst(r.content)
    assert "Kuil" in tekst and "Geen beeld bewaard" in tekst


@pytest.mark.parametrize("ext", ["pdf", "xlsx"])
def test_rapport_van_andere_organisatie_is_onvindbaar(client, admin_user, ext):
    db = SessionLocal()
    try:
        vreemd = Schouwrit(organization_id="een-andere-org", gebied="Elders", status="bezig",
                           privacy_modus="gericht", created_by=admin_user.id)
        db.add(vreemd)
        db.commit()
        vreemd_id = vreemd.id
    finally:
        db.close()
    r = client.get(f"/api/schouw/ritten/{vreemd_id}/rapport.{ext}", headers=auth(admin_user))
    assert r.status_code == 404


def test_rode_vak_komt_op_het_beeld():
    from PIL import Image
    data, verhouding = schouw_rapport.beeld_met_vak(_jpeg(), [0.25, 0.25, 0.75, 0.75], "Kuilen")
    assert verhouding > 0
    beeld = Image.open(io.BytesIO(data)).convert("RGB")
    w, h = beeld.size
    # Binnen het vak is het roder dan erbuiten.
    binnen = beeld.getpixel((w // 2, h // 2))
    buiten = beeld.getpixel((2, 2))
    assert binnen[0] - binnen[1] > buiten[0] - buiten[1]
    assert schouw_rapport.beeld_met_vak(b"geen beeld", None) is None


@pytest.mark.parametrize('ext', ['pdf', 'xlsx'])
def test_rapport_vraagt_authenticatie_en_actieve_module(client, admin_user, ext):
    rit_id = _rit(client, admin_user)
    url = f'/api/schouw/ritten/{rit_id}/rapport.{ext}'
    assert client.get(url).status_code == 401
    db = SessionLocal()
    try:
        from models import Organization
        org = db.get(Organization, admin_user.organization_id)
        org.enabled_modules = '[]'
        db.commit()
    finally:
        db.close()
    assert client.get(url, headers=auth(admin_user)).status_code == 403


def test_excel_schadegegevens_en_opmaak(client, admin_user, model):
    from openpyxl import load_workbook
    rit_id = _rit(client, admin_user)
    model['antwoorden'].append({'bruikbaar': True, 'wegschade': [dict(KUIL)]})
    _opname(client, admin_user, rit_id, 1)
    schade = _schades(rit_id)[0]
    client.patch(f'/api/schouw/waarnemingen/{schade.id}', headers=auth(admin_user),
                 json={'afgewezen': True, 'afwijs_reden': 'schaduw'})
    response = client.get(f'/api/schouw/ritten/{rit_id}/rapport.xlsx', headers=auth(admin_user))
    assert response.status_code == 200
    import zipfile
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        assert any(n.startswith('xl/media/') for n in archive.namelist())
        assert not any(n.startswith('xl/externalLinks/') for n in archive.namelist())
    ws = load_workbook(io.BytesIO(response.content))['Schades']
    assert 'schaduw' in ws['I12'].value.lower()
    assert ws['E12'].value == KUIL['ernst']
    assert ws['H12'].number_format == '0%'
    image = ws._images[0]
    assert image.anchor._from.col == 1 and image.anchor._from.row == 11
    # Afbeelding blijft binnen de fotokolom en de rij, ook in afdruk.
    assert image.anchor.ext.cx / 9525 <= ws.column_dimensions['B'].width * 7
    assert image.anchor.ext.cy / 9525 <= ws.row_dimensions[12].height / 0.75


def test_export_routes_zijn_uniek():
    from routers.schouw_router import router
    for ext in ("pdf", "xlsx"):
        routes = [r for r in router.routes
                  if getattr(r, "path", "") == f"/api/schouw/ritten/{{rit_id}}/rapport.{ext}"
                  and "GET" in getattr(r, "methods", set())]
        assert len(routes) == 1
